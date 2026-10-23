"""Bounded native source facts, separate from the adapter's display projection.

Streams are single use. Limits constrain parser work, not a claim about unseen
source totals. Structured scalars retain native types; no field name supplies
semantics and strings containing another format are never decoded.
"""

from collections.abc import Callable, Iterator
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import TYPE_CHECKING

from sage.adapters.interfaces import HEADING_PATH_SEPARATOR

if TYPE_CHECKING:
    from sage.source_adapters.base import ProjectionResult


@dataclass(frozen=True)
class SourceFact:
    region: str
    field_path: str
    location: str
    order: int
    record: str
    value: object
    rendering: str
    context: str = ""
    authored_labels: tuple[str, ...] = ()


@dataclass(frozen=True)
class FactLimits:
    max_source_bytes: int = 64 * 1024 * 1024
    max_depth: int = 64
    max_parse_nodes: int = 400_000
    max_rows: int = 100_000
    max_columns: int = 256

    @classmethod
    def hard_limits(cls) -> dict[str, int]:
        """Single policy authority for native parser hard maxima."""
        return {
            "max_source_bytes": 256 * 1024 * 1024,
            "max_depth": 128,
            "max_parse_nodes": 1_600_000,
            "max_rows": 400_000,
            "max_columns": 1024,
        }

    def __post_init__(self) -> None:
        maxima = self.hard_limits()
        for key, value in vars(self).items():
            if type(value) is not int or value < 1:
                raise ValueError("fact limits must be positive integers")
            if value > maxima[key]:
                raise ValueError(f"{key} exceeds hard maximum {maxima[key]}")


class FactStream:
    """Single-pass facts and known limitations, settled only when consumed."""

    def __init__(self, factory: Callable[[], Iterator[SourceFact]], limits: FactLimits) -> None:
        self._limits = limits
        self.limitations: list[str] = []
        self.exhausted = False
        self._iterator = factory()

    @property
    def limits(self) -> FactLimits:
        """The immutable parser limits used for this stream."""
        return self._limits

    def __iter__(self) -> Iterator[SourceFact]:
        return self._iterator


class _BoundReached(ValueError):
    pass


def _render(value: object) -> str:
    import json

    if isinstance(value, str):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return json.dumps(value, ensure_ascii=False, default=str)


def source_facts(
    path: Path,
    projection: "ProjectionResult | None",
    limits: FactLimits,
    *,
    native: bool | None = None,
) -> FactStream:
    """Inspect native data lazily; stop without reading a remainder to count it."""
    stream: FactStream

    def generate() -> Iterator[SourceFact]:
        try:
            if native is not False and path.suffix.lower() in {
                ".json",
                ".jsonl",
                ".yaml",
                ".yml",
                ".toml",
                ".xml",
                ".csv",
                ".tsv",
                ".xlsx",
            }:
                if path.stat().st_size > limits.max_source_bytes:
                    raise _BoundReached("source byte admission limit reached; source not inspected")
                for fact in _native(path, limits, stream.limitations):
                    if isinstance(fact.value, str):
                        stripped = fact.value.strip()
                        if (stripped.startswith("{") and stripped.endswith("}")) or (
                            stripped.startswith("[") and stripped.endswith("]")
                        ):
                            limitation = (
                                "container-like text retained without decoding; "
                                "embedded structure unresolved"
                            )
                            if limitation not in stream.limitations:
                                stream.limitations.append(limitation)
                    yield fact
            elif projection is not None:
                yield from _narrative(projection, limits)
            else:
                stream.limitations.append("no native facts reader or narrative projection")
                return
        except _BoundReached as exc:
            stream.limitations.append(str(exc))
            return
        except (ValueError, OSError, RecursionError) as exc:
            from sage.source_adapters.base import SourceReadError

            raise SourceReadError(f"Cannot inspect source {path}: {exc}") from exc
        stream.exhausted = True

    stream = FactStream(generate, limits)
    return stream


def _narrative(projection: "ProjectionResult", limits: FactLimits) -> Iterator[SourceFact]:
    bodies = ((h.path, h.content) for h in projection.headings)
    if not projection.headings:
        bodies = iter((("document", projection.text),))
    elif projection.preamble:
        from itertools import chain

        bodies = chain((("preamble", projection.preamble),), bodies)
    order = 0
    for region, text in bodies:
        if len(region) > 500:
            raise _BoundReached("narrative location length limit; remainder unknown")
        # Split incrementally, avoiding a parallel list of the whole projection.
        import re

        for match in re.finditer(r"\S[\s\S]*?(?=\n\s*\n|\Z)", text):
            for segment_start in range(match.start(), match.end(), 16_384):
                if order >= limits.max_parse_nodes:
                    raise _BoundReached("narrative fact limit reached; remainder unknown")
                value = text[segment_start : min(segment_start + 16_384, match.end())]
                yield SourceFact(
                    region,
                    "text",
                    f"{region}:char:{segment_start}",
                    order,
                    region,
                    value,
                    value,
                    region if region != "document" else "",
                )
                order += 1


def _walk(
    value: object, limits: FactLimits, prefix: str = "$", start: int = 0
) -> Iterator[SourceFact]:
    order = start
    nodes = 0

    def visit(
        item: object, path: str, field: str, depth: int, record: str, labels: tuple[str, ...]
    ) -> Iterator[SourceFact]:
        nonlocal nodes, order
        nodes += 1
        if nodes > limits.max_parse_nodes or depth > limits.max_depth:
            raise _BoundReached("structured traversal limit reached; remainder unknown")
        if isinstance(item, dict):
            for key, child in item.items():
                import json

                label = key if isinstance(key, str) else str(key)
                if len(label) > 500:
                    raise _BoundReached("structured location length limit; remainder unknown")
                segment = f"[{json.dumps(label, ensure_ascii=False)}]"
                if max(len(path), len(field)) + len(segment) > 512:
                    raise _BoundReached("structured location length limit; remainder unknown")
                yield from visit(
                    child,
                    path + segment,
                    field + segment,
                    depth + 1,
                    record,
                    labels + (label,) if isinstance(key, str) else labels,
                )
        elif isinstance(item, list):
            for index, child in enumerate(item):
                if max(len(path), len(field)) + len(str(index)) + 2 > 512:
                    raise _BoundReached("structured location length limit; remainder unknown")
                child_path = f"{path}[{index}]"
                yield from visit(child, child_path, field + "[]", depth + 1, child_path, labels)
        else:
            yield SourceFact(
                "document",
                field,
                path,
                order,
                record,
                item,
                _render(item),
                HEADING_PATH_SEPARATOR.join(labels),
                labels,
            )
            order += 1

    yield from visit(value, prefix, "$", 0, prefix, ())


def _json_guard(text: str, limits: FactLimits) -> int:
    """Lexical work guard before allocating a JSON tree, ignoring string syntax."""
    quoted = escaped = False
    depth = nodes = 0
    for char in text:
        if quoted:
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == '"':
                quoted = False
            continue
        if char == '"':
            quoted = True
            nodes += 1
        elif char in "[{":
            depth += 1
            nodes += 1
        elif char in "]}":
            depth -= 1
        elif char in ",:":
            nodes += 1
        if depth > limits.max_depth or nodes > limits.max_parse_nodes:
            raise _BoundReached("JSON parser depth/node limit reached before tree allocation")

    return max(nodes, 1)


def _native(path: Path, limits: FactLimits, limitations: list[str]) -> Iterator[SourceFact]:
    import csv
    import json
    import tomllib

    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        yield from _xlsx(path, limits)
        return
    if suffix in {".csv", ".tsv"}:
        with path.open(encoding="utf-8-sig", newline="") as source:
            reader = csv.reader(source, delimiter="\t" if suffix == ".tsv" else ",", strict=True)
            order = 0
            for row_index in range(limits.max_rows):
                line = reader.line_num + 1
                row = next(reader, None)
                if row is None:
                    return
                if len(row) > limits.max_columns:
                    raise _BoundReached("delimited column limit reached; remainder unknown")
                for column, value in enumerate(row, 1):
                    if order >= limits.max_parse_nodes:
                        raise _BoundReached("delimited fact limit reached; remainder unknown")
                    yield SourceFact(
                        "document",
                        f"column:{column}",
                        f"line:{line}:column:{column}",
                        order,
                        f"row:{row_index + 1}",
                        value,
                        value,
                    )
                    order += 1
            raise _BoundReached("delimited row limit reached; remainder unknown")
    if suffix == ".jsonl":
        with path.open(encoding="utf-8-sig") as source:
            order = 0
            for index in range(limits.max_rows):
                line = source.readline()
                if not line:
                    return
                if not line.strip():
                    continue
                _json_guard(line, limits)
                for fact in _walk(json.loads(line), limits, f"line:{index + 1}", order):
                    if order >= limits.max_parse_nodes:
                        raise _BoundReached("JSONL fact limit reached; remainder unknown")
                    yield fact
                    order += 1
            raise _BoundReached("JSONL row limit reached; remainder unknown")
    scalar_spans: list[tuple[int, int]] = []
    text = path.read_text(encoding="utf-8-sig")
    if suffix == ".xml":
        yield from _xml(text, limits)
        return
    if suffix == ".json":
        _json_guard(text, limits)
        data = json.loads(text)
    else:
        import yaml

        # YAML's event parser can guard all YAML/TOML nesting before native
        # construction. TOML is additionally constrained by its lexical table
        # path depth; repeated dotted keys can allocate nested native tables.
        depth = nodes = 0
        aliases = False
        if suffix in {".yaml", ".yml"}:
            for event in yaml.parse(text):
                nodes += 1
                if isinstance(event, yaml.ScalarEvent):
                    scalar_spans.append((event.start_mark.index, event.end_mark.index))
                if isinstance(event, (yaml.MappingStartEvent, yaml.SequenceStartEvent)):
                    depth += 1
                elif isinstance(event, (yaml.MappingEndEvent, yaml.SequenceEndEvent)):
                    depth -= 1
                aliases |= isinstance(event, yaml.AliasEvent)
                if depth > limits.max_depth or nodes > limits.max_parse_nodes:
                    raise _BoundReached(
                        "YAML parser depth/node limit reached before tree allocation"
                    )
            if aliases:
                limitations.append(
                    "YAML aliases retained lexically; expanded native values uninspected"
                )
                for order, line in enumerate(text.splitlines()):
                    if order >= limits.max_parse_nodes:
                        raise _BoundReached("YAML lexical line limit reached; remainder unknown")
                    yield SourceFact(
                        "document", "$source", f"line:{order + 1}", order, "document", line, line
                    )
                return
            documents = list(yaml.safe_load_all(text))
            data = documents[0] if len(documents) == 1 else documents
        else:
            # Ignore quoted strings and comments while bounding TOML's parser.
            scalar_spans = _toml_guard(text, limits)
            data = tomllib.loads(text)
    order = 0
    for fact in _walk(data, limits):
        yield fact
        order = fact.order + 1
    if suffix in {".yaml", ".yml", ".toml"}:
        # Comments are a supplemental lexical region; native record ordering
        # remains intact and source-line locations preserve their authored place.
        for line, comment in _comments(text, scalar_spans):
            if order >= limits.max_parse_nodes:
                raise _BoundReached("comment fact limit reached; remainder unknown")
            yield SourceFact(
                "comments",
                "$comment",
                f"line:{line}:comment",
                order,
                f"line:{line}",
                comment,
                comment,
            )
            order += 1


def _comments(text: str, spans: list[tuple[int, int]]) -> Iterator[tuple[int, str]]:
    """Only hashes outside parser-identified scalar/string spans are comments."""
    import re

    protected = iter(spans)
    span = next(protected, None)
    last = 0
    comment_end = -1
    line = 1
    for match in re.finditer(r"#", text):
        if match.start() < comment_end:
            continue
        line += text.count("\n", last, match.start())
        last = match.start()
        while span is not None and span[1] <= match.start():
            span = next(protected, None)
        if span is None or match.start() < span[0]:
            end = text.find("\n", match.start())
            comment_end = end if end >= 0 else len(text)
            yield line, text[match.start() : comment_end].rstrip("\r")
            # Hashes later on this same comment line belong to this comment.
            while span is not None and span[1] <= match.start():
                span = next(protected, None)


def _xlsx(path: Path, limits: FactLimits) -> Iterator[SourceFact]:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    from sage.source_adapters.base import check_zip_package

    check_zip_package(path)
    columns_clipped = _xlsx_guard(path, limits)
    workbook = load_workbook(path, read_only=True, data_only=False)
    order = rows = 0
    try:
        for sheet in workbook:
            for row in sheet.iter_rows(max_row=limits.max_rows, max_col=limits.max_columns):
                if rows >= limits.max_rows:
                    raise _BoundReached("workbook row limit reached; remainder unknown")
                rows += 1
                for column, cell in enumerate(row, 1):
                    if cell.value is None:
                        continue
                    if order >= limits.max_parse_nodes:
                        raise _BoundReached("workbook fact limit reached; remainder unknown")
                    address = f"{sheet.title}!{get_column_letter(column)}{cell.row}"
                    yield SourceFact(
                        sheet.title,
                        f"column:{column}",
                        address,
                        order,
                        f"{sheet.title}!row:{cell.row}",
                        cell.value,
                        _render(cell.value),
                        sheet.title,
                        (sheet.title,),
                    )
                    order += 1
            if rows >= limits.max_rows:
                raise _BoundReached("workbook row limit reached; remainder unknown")
        columns_clipped |= any((s.max_column or 0) > limits.max_columns for s in workbook)
        if columns_clipped:
            raise _BoundReached("workbook column limit reached; remainder unknown")
    finally:
        workbook.close()


def _xml(text: str, limits: FactLimits) -> Iterator[SourceFact]:
    """SAX callbacks bound work before constructing any native XML tree."""
    from collections import deque
    from xml.sax.handler import ContentHandler, LexicalHandler, property_lexical_handler

    from defusedxml.expatreader import create_parser

    queue: deque[SourceFact] = deque()

    class Handler(ContentHandler, LexicalHandler):
        def __init__(self) -> None:
            self.stack: list[tuple[str, dict[str, int]]] = []
            self.roots: dict[str, int] = {}
            self.order = self.nodes = 0

        def emit(self, field: str, location: str, value: str) -> None:
            self.nodes += 1
            if self.nodes > limits.max_parse_nodes:
                raise _BoundReached("XML parser node limit reached; remainder unknown")
            record = self.stack[-1][0] if self.stack else "document"
            queue.append(SourceFact("document", field, location, self.order, record, value, value))
            self.order += 1

        def startElement(self, name: str, attrs: object) -> None:  # noqa: N802
            self.nodes += 1
            if len(self.stack) >= limits.max_depth or self.nodes > limits.max_parse_nodes:
                raise _BoundReached("XML parser depth/node limit reached; remainder unknown")
            parent = self.stack[-1][0] if self.stack else ""
            if len(name) + len(parent) + 16 > 512:
                raise _BoundReached("XML location length limit; remainder unknown")
            siblings = self.stack[-1][1] if self.stack else self.roots
            siblings[name] = siblings.get(name, 0) + 1
            path = (self.stack[-1][0] if self.stack else "") + f"/{name}[{siblings[name]}]"
            self.stack.append((path, {}))
            for key, value in attrs.items():  # type: ignore[attr-defined]
                if len(key) + len(path) + 2 > 512:
                    raise _BoundReached("XML attribute location length limit; remainder unknown")
                self.emit(f"{name}/@{key}", f"{path}/@{key}", value)

        def endElement(self, name: str) -> None:  # noqa: N802
            self.stack.pop()

        def comment(self, content: str) -> None:
            self.emit("$comment", f"comment:{self.order}", content)

        def processingInstruction(self, target: str, data: str) -> None:  # noqa: N802
            if len(target) > 500:
                raise _BoundReached("XML instruction location length limit; remainder unknown")
            self.emit(target, f"instruction:{self.order}", data)

        def characters(self, content: str) -> None:
            if content and self.stack:
                self.emit("text()", f"{self.stack[-1][0]}/text():{self.order}", content)

    parser = create_parser(forbid_dtd=False, forbid_entities=True, forbid_external=True)
    handler = Handler()
    parser.setContentHandler(handler)
    parser.setProperty(property_lexical_handler, handler)
    for offset in range(0, len(text), 4096):
        parser.feed(text[offset : offset + 4096])
        while queue:
            yield queue.popleft()
    parser.close()
    yield from queue


def _xlsx_guard(path: Path, limits: FactLimits, *, projection: bool = False) -> bool:
    """Bound XML work before openpyxl allocates shared strings/styles/rows."""
    import zipfile
    from xml.sax.handler import ContentHandler

    from defusedxml.expatreader import create_parser

    class Guard(ContentHandler):
        def __init__(self) -> None:
            self.nodes = self.depth = 0
            self.columns_clipped = False
            self.rows = 0

        def startElement(self, name: str, attrs: object) -> None:  # noqa: N802
            self.depth += 1
            self.nodes += 1
            local_name = name.rsplit(":", 1)[-1]
            if local_name == "row" and projection:
                self.rows += 1
                reference = attrs.get("r", "0")  # type: ignore[attr-defined]
                if len(reference) > 9 or not reference.isdigit():
                    raise _BoundReached("workbook row address limit")
                if self.rows > limits.max_rows or int(reference) > limits.max_rows:
                    raise _BoundReached("workbook row limit")
            if local_name == "c":
                reference = attrs.get("r", "")  # type: ignore[attr-defined]
                if len(reference) > 16:
                    raise _BoundReached("workbook cell address length limit")
                column = 0
                for char in reference:
                    if not "A" <= char <= "Z":
                        break
                    column = column * 26 + ord(char) - ord("A") + 1
                self.columns_clipped |= column > limits.max_columns
            if self.depth > limits.max_depth or self.nodes > limits.max_parse_nodes:
                raise _BoundReached("workbook XML depth/node limit before native allocation")

        def endElement(self, name: str) -> None:  # noqa: N802
            self.depth -= 1

    handler = Guard()
    with zipfile.ZipFile(path) as archive:
        for member in archive.infolist():
            if not member.filename.endswith((".xml", ".rels")):
                continue
            parser = create_parser(forbid_dtd=True, forbid_entities=True, forbid_external=True)
            parser.setContentHandler(handler)
            with archive.open(member) as source:
                while block := source.read(4096):
                    parser.feed(block)
            parser.close()
    return handler.columns_clipped


def _toml_guard(text: str, limits: FactLimits) -> list[tuple[int, int]]:
    """A lexical guard for TOML, including dotted paths and multiline strings."""
    spans: list[tuple[int, int]] = []
    index = depth = nodes = dots = 0
    in_key = True
    while index < len(text):
        char = text[index]
        if char == "#":
            end = text.find("\n", index)
            index = len(text) if end < 0 else end
            continue
        if char in {'"', "'"}:
            start = index
            marker = char * 3 if text.startswith(char * 3, index) else char
            index += len(marker)
            while index < len(text):
                if char == '"' and text[index] == "\\":
                    index += 2
                elif text.startswith(marker, index):
                    index += len(marker)
                    break
                else:
                    index += 1
            spans.append((start, index))
            nodes += 1
        else:
            if char in "[{":
                depth += 1
            elif char in "]}":
                depth = max(0, depth - 1)
            elif char == "." and in_key:
                dots += 1
            elif char == "=":
                in_key = False
            elif char == "\n":
                dots = 0
                in_key = True
            if not char.isspace():
                nodes += 1
            index += 1
        if depth > limits.max_depth or dots >= limits.max_depth or nodes > limits.max_parse_nodes:
            raise _BoundReached("TOML parser depth/node limit reached before tree allocation")
    return spans


def check_projection_work(path: Path, limits: FactLimits) -> None:
    """Refuse excess parser work before an adapter materializes its projection.

    This admission boundary intentionally raises a source error rather than
    presenting a partial inspection as permission for an unbounded projector.
    JSONL node work is aggregated across lines. YAML aliases are admitted only
    after their finite expanded structural cost has been established.
    """
    import csv
    import io
    import zipfile
    from xml.sax import SAXException

    import yaml
    from defusedxml import DefusedXmlException

    from sage.source_adapters.base import SourceReadError, check_zip_package

    suffix = path.suffix.lower()
    if suffix not in {".json", ".jsonl", ".yaml", ".yml", ".toml", ".xml", ".xlsx", ".csv", ".tsv"}:
        return
    try:
        if path.stat().st_size > limits.max_source_bytes:
            raise _BoundReached("source byte admission limit")
        if suffix == ".xlsx":
            check_zip_package(path)
            # Projection already clamps row width independently. Wide sources
            # remain readable through that existing preview; facts report their
            # separately configured inspection width as limited.
            _xlsx_guard(path, limits, projection=True)
            return
        # Admission precedes this whole-source allocation. The bounded read
        # also refuses growth between stat and open instead of trusting stat.
        with path.open("rb") as source:
            raw = source.read(limits.max_source_bytes + 1)
        if len(raw) > limits.max_source_bytes:
            raise _BoundReached("source byte admission limit")
        text = raw.decode("utf-8-sig")
        if suffix == ".json":
            _json_guard(text, limits)
        elif suffix == ".jsonl":
            nodes = rows = 0
            for line in io.StringIO(text):
                rows += 1
                if rows > limits.max_rows:
                    raise _BoundReached("JSONL row limit")
                if line.strip():
                    nodes += _json_guard(line, limits)
                    if nodes > limits.max_parse_nodes:
                        raise _BoundReached("aggregate JSONL parser node limit")
        elif suffix in {".yaml", ".yml"}:
            _yaml_projection_guard(text, limits)
        elif suffix == ".toml":
            _toml_guard(text, limits)
        elif suffix == ".xml":
            for _fact in _xml(text, limits):
                pass
        else:
            reader = csv.reader(
                io.StringIO(text, newline=""),
                delimiter="\t" if suffix == ".tsv" else ",",
                strict=True,
            )
            nodes = 0
            for row_index, row in enumerate(reader, 1):
                nodes += len(row)
                if (
                    row_index > limits.max_rows
                    or len(row) > limits.max_columns
                    or nodes > limits.max_parse_nodes
                ):
                    raise _BoundReached("delimited row/column/node limit")
    except (
        ValueError,
        OSError,
        RecursionError,
        yaml.YAMLError,
        SAXException,
        DefusedXmlException,
        csv.Error,
        zipfile.BadZipFile,
    ) as exc:
        raise SourceReadError(f"Source {path} exceeds safe projection work: {exc}") from exc


def _yaml_projection_guard(text: str, limits: FactLimits) -> None:
    """Bound finite alias expansion before SafeLoader's merge flattening."""
    import yaml

    # A frame is [expanded cost, anchor]; no source subtree is copied here.
    stack: list[list] = []
    anchors: dict[str, int] = {}
    nodes = total = 0

    def add(cost: int) -> None:
        nonlocal total
        if stack:
            if cost > limits.max_parse_nodes - stack[-1][0]:
                raise _BoundReached("YAML expanded alias/node limit")
            stack[-1][0] += cost
        else:
            if cost > limits.max_parse_nodes - total:
                raise _BoundReached("YAML expanded document/node limit")
            total += cost

    for event in yaml.parse(text):
        nodes += 1
        if nodes > limits.max_parse_nodes:
            raise _BoundReached("YAML parser event/node limit")
        if isinstance(event, yaml.DocumentStartEvent):
            anchors.clear()
        elif isinstance(event, (yaml.MappingStartEvent, yaml.SequenceStartEvent)):
            if len(stack) >= limits.max_depth:
                raise _BoundReached("YAML parser depth limit")
            stack.append([1, event.anchor])
        elif isinstance(event, (yaml.MappingEndEvent, yaml.SequenceEndEvent)):
            cost, anchor = stack.pop()
            if anchor:
                anchors[anchor] = cost
            add(cost)
        elif isinstance(event, yaml.ScalarEvent):
            if event.anchor:
                anchors[event.anchor] = 1
            add(1)
        elif isinstance(event, yaml.AliasEvent):
            if event.anchor not in anchors:
                raise _BoundReached("YAML recursive/incomplete alias cannot be safely expanded")
            add(anchors[event.anchor])
