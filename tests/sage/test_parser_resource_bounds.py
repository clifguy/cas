"""Source parsing is bounded in memory and time, and runs off the event loop.

A small input must not cost the process memory or time out of proportion to
its size: a workbook's declared dimension does not size its rows, an Office
package is checked against fixed decompression limits before any library
inflates it, the structured-data separation pass works within a fixed parse
budget, and the filename parser does linear work on a bounded input.
Projection runs on a worker thread, so a slow source does not stall the
server's other requests.
"""

from __future__ import annotations

import asyncio
import re
import threading
import time
import tracemalloc
import zipfile
from pathlib import Path

import pytest
from hypothesis import example, given, settings
from hypothesis import strategies as st

from sage.services import filename_parser
from sage.services.filename_parser import FilenameParser
from sage.source_adapters import base as adapter_base
from sage.source_adapters.base import SourceReadError, check_zip_package
from sage.source_adapters.docx_adapter import DocxAdapter
from sage.source_adapters.pptx_adapter import PptxAdapter
from sage.source_adapters.structured_data_adapter import StructuredDataAdapter
from sage.source_adapters.xlsx_adapter import XlsxAdapter

# ---------------------------------------------------------------------------
# Fixture builders
# ---------------------------------------------------------------------------


def _rewrite_member(path: Path, member: str, transform) -> None:
    """Rewrite one member of a zip package in place, keeping the others."""
    with zipfile.ZipFile(path) as zin:
        items = [(info, zin.read(info.filename)) for info in zin.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info, data in items:
            if info.filename == member:
                data = transform(data)
            zout.writestr(info, data)


def _add_member(path: Path, name: str, data: bytes) -> None:
    with zipfile.ZipFile(path, "a", zipfile.ZIP_DEFLATED) as zout:
        zout.writestr(name, data)


def _dimension_inflated_xlsx(tmp_path: Path, rows: int) -> Path:
    """A workbook of ``rows`` one-cell rows whose sheet declares the maximum dimension."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    for i in range(rows):
        ws.append([f"r{i}"])
    path = tmp_path / "inflated.xlsx"
    wb.save(str(path))
    wb.close()
    _rewrite_member(
        path,
        "xl/worksheets/sheet1.xml",
        lambda xml: re.sub(rb'<dimension ref="[^"]*"/>', b'<dimension ref="A1:XFD1048576"/>', xml),
    )
    return path


def _docx(tmp_path: Path) -> Path:
    import docx

    doc = docx.Document()
    doc.add_paragraph("Body text.")
    path = tmp_path / "doc.docx"
    doc.save(str(path))
    return path


def _pptx(tmp_path: Path) -> Path:
    from pptx import Presentation

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "Title"
    path = tmp_path / "deck.pptx"
    prs.save(str(path))
    return path


def _xlsx(tmp_path: Path) -> Path:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.append(["a", "b"])
    path = tmp_path / "book.xlsx"
    wb.save(str(path))
    wb.close()
    return path


_PACKAGES = {
    "docx": (_docx, DocxAdapter),
    "pptx": (_pptx, PptxAdapter),
    "xlsx": (_xlsx, XlsxAdapter),
}


# ---------------------------------------------------------------------------
# xlsx: rows are not sized by the declared dimension
# ---------------------------------------------------------------------------


async def test_xlsx_declared_dimension_does_not_size_rows(tmp_path):
    """TEST-SAGE-AD-222: a workbook declaring the maximum sheet dimension
    over a few thousand one-cell rows projects within a fixed memory budget
    and reports the columns that hold data.

    Each row padded to the declared 16,384 columns costs about 130 KB, so the
    unbounded read peaks in the hundreds of megabytes on this fixture; the
    reported column count of 1 also fails if the dimension sizes the rows.
    """
    path = _dimension_inflated_xlsx(tmp_path, rows=1500)

    tracemalloc.start()
    try:
        result = await XlsxAdapter().project(path)
        _current, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()

    assert peak < 64 * 1024 * 1024, f"peak {peak / 2**20:.0f} MiB"
    assert result.metadata["dimensions"]["Data"] == {"rows": 1500, "columns": 1}
    assert "A1 [string]\nr0" in result.text
    assert "A1500 [string]\nr1499" in result.text
    # Counted to the last row holding a value, not scanned out to the
    # declared extent and reported as a lower bound.
    assert len(result.passages) == 1500
    assert result.projection_complete is True


async def test_xlsx_row_width_is_capped(tmp_path):
    """TEST-SAGE-AD-234: rows holding a value in the sheet's last column are
    read sparsely, preserving far authored cells within a stored-cell limit.

    The coordinate envelope is not a rectangular allocation. An explicitly
    smaller stored-cell budget refuses the source instead of dropping data.
    """
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Wide"
    for i in range(1, 51):
        ws.cell(row=i, column=1, value=f"r{i}")
        ws.cell(row=i, column=16384, value="edge")
    path = tmp_path / "wide.xlsx"
    wb.save(str(path))
    wb.close()

    result = await XlsxAdapter().project(path)

    columns = result.metadata["dimensions"]["Wide"]["columns"]
    assert columns == 16384
    assert len(result.passages) == 100
    assert "XFD50 [string]\nedge" in result.text
    with pytest.raises(SourceReadError, match="stored column limit"):
        await XlsxAdapter().project(path, {"discovery_facts": {"max_columns": 1}})


async def test_xlsx_row_scan_is_capped(tmp_path):
    """TEST-SAGE-AD-223: stored-row work is admitted explicitly; an excess
    refuses projection rather than advertising a truncated preview as content.
    """
    path = _dimension_inflated_xlsx(tmp_path, rows=150)

    result = await XlsxAdapter().project(path)

    assert result.metadata["dimensions"]["Data"]["rows"] == 150
    assert "A150 [string]\nr149" in result.text
    with pytest.raises(SourceReadError, match="row limit"):
        await XlsxAdapter().project(path, {"discovery_facts": {"max_rows": 100}})


# ---------------------------------------------------------------------------
# xlsx: only cells inside worksheet rows count toward the stored-column bound
# ---------------------------------------------------------------------------


def _rows_xlsx(tmp_path: Path, rows: list[list[object]], title: str = "Data") -> Path:
    """A one-sheet workbook whose rows hold exactly the given values."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for values in rows:
        ws.append(values)
    path = tmp_path / f"{title}.xlsx"
    wb.save(str(path))
    wb.close()
    return path


def _with_calc_chain(path: Path, entries: int) -> None:
    """Add a calculation chain of ``entries`` formula-cell references.

    The chain is a flat list of ``<c>`` elements with no rows, as Excel
    writes it for every workbook holding formulas.
    """
    cells = "".join(f'<c r="B{i}" i="1"/>' for i in range(1, entries + 1))
    _add_member(
        path,
        "xl/calcChain.xml",
        (
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            f"{cells}</calcChain>"
        ).encode(),
    )
    _rewrite_member(
        path,
        "[Content_Types].xml",
        lambda xml: xml.replace(
            b"</Types>",
            b'<Override PartName="/xl/calcChain.xml" ContentType="application/'
            b'vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/></Types>',
        ),
    )


def _with_column_span(path: Path, first: int) -> None:
    """Declare column formatting from column ``first`` out to XFD."""
    _rewrite_member(
        path,
        "xl/worksheets/sheet1.xml",
        lambda xml: xml.replace(
            b"<sheetData>",
            f'<cols><col min="{first}" max="16384" width="9" customWidth="1"/></cols>'
            "<sheetData>".encode(),
            1,
        ),
    )


def _guard_detail(exc: pytest.ExceptionInfo[SourceReadError]) -> str:
    return str(exc.value).split("exceeds safe projection work: ", 1)[1]


async def test_xlsx_calc_chain_does_not_count_as_stored_columns(tmp_path):
    """TEST-SAGE-AD-238: a narrow workbook with more formula cells than the
    stored-column bound, and column formatting declared out to XFD, is
    admitted and projects in full.

    The calculation chain lists one ``<c>`` per formula cell outside any row;
    counting those as cells of the last worksheet row refuses the workbook.
    """
    from sage.source_adapters.structure import FactLimits, check_projection_work

    path = _rows_xlsx(tmp_path, [[i, f"=A{i}*2"] for i in range(1, 301)])
    _with_calc_chain(path, 300)
    _with_column_span(path, 3)

    check_projection_work(path, FactLimits())
    result = await XlsxAdapter().project(path)

    assert "A300 [" in result.text
    assert "B300 [" in result.text


async def test_xlsx_column_formatting_span_is_not_a_stored_column(tmp_path):
    """TEST-SAGE-AD-239: a column-formatting declaration spanning to XFD
    costs nothing to project and does not count toward the stored-column
    bound; two stored cells per row are admitted at a bound of two.
    """
    path = _rows_xlsx(tmp_path, [["a", "b"]] * 5)
    _with_column_span(path, 1)

    result = await XlsxAdapter().project(path, {"discovery_facts": {"max_columns": 2}})

    assert "B5 [string]\nb" in result.text


async def test_xlsx_row_count_does_not_leak_into_a_rowless_part(tmp_path):
    """TEST-SAGE-AD-240: a last worksheet row holding exactly the bound is
    admitted even when a later rowless part carries a ``<c>`` element, while
    one more stored cell in that row is refused.
    """
    at_bound = _rows_xlsx(tmp_path, [["a", "b"]], title="AtBound")
    _with_calc_chain(at_bound, 1)
    over_bound = _rows_xlsx(tmp_path, [["a", "b", "c"]], title="OverBound")
    limits = {"discovery_facts": {"max_columns": 2}}

    result = await XlsxAdapter().project(at_bound, limits)

    assert "B1 [string]\nb" in result.text
    with pytest.raises(SourceReadError, match="stored column limit"):
        await XlsxAdapter().project(over_bound, limits)


@pytest.mark.parametrize("target", ["/xl/worksheets/sheet1.xml", "worksheets/sheet1.xml"])
async def test_xlsx_column_refusal_names_sheet_part_row_and_limit(tmp_path, target):
    """TEST-SAGE-AD-241: a row genuinely over the stored-column bound is
    refused with the sheet, the package part, the row, the measured cell
    count and the limit, so the refusal can be acted on without unpacking
    the workbook. The sheet resolves whether the workbook relationship
    names its part absolutely or relative to the workbook part.
    """
    # The empty second row is not stored, so the sheet's row number (3)
    # differs from the count of rows read (2).
    path = _rows_xlsx(tmp_path, [["a", "b"], [], ["a", "b", "c"]], title="Wide")
    _rewrite_member(
        path,
        "xl/_rels/workbook.xml.rels",
        lambda xml: xml.replace(
            b'Target="/xl/worksheets/sheet1.xml"', f'Target="{target}"'.encode()
        ),
    )

    with pytest.raises(SourceReadError) as exc:
        await XlsxAdapter().project(path, {"discovery_facts": {"max_columns": 2}})

    assert _guard_detail(exc) == (
        "workbook stored column limit: sheet 'Wide' in part "
        "'xl/worksheets/sheet1.xml' row 3 holds 3 stored cells; max_columns is 2"
    )


async def test_xlsx_row_refusal_names_count_and_limit(tmp_path):
    """TEST-SAGE-AD-242: a workbook over the stored-row bound is refused
    with the measured row count and the limit.
    """
    path = _dimension_inflated_xlsx(tmp_path, rows=150)

    with pytest.raises(SourceReadError) as exc:
        await XlsxAdapter().project(path, {"discovery_facts": {"max_rows": 100}})

    assert _guard_detail(exc) == "workbook row limit: 101 rows across the workbook; max_rows is 100"


async def test_xlsx_refusal_renders_a_hostile_sheet_name_inertly(tmp_path):
    """TEST-SAGE-AD-243: a sheet name is document-derived, so a refusal
    quoting it stays on one line and within a fixed length whatever the
    workbook declares.
    """
    path = _rows_xlsx(tmp_path, [["a", "b"]])
    _rewrite_member(
        path,
        "xl/workbook.xml",
        lambda xml: xml.replace(b'name="Data"', b'name="&#10;' + b"X" * 200 + b'"'),
    )

    with pytest.raises(SourceReadError) as exc:
        await XlsxAdapter().project(path, {"discovery_facts": {"max_columns": 1}})

    detail = _guard_detail(exc)
    assert "\n" not in detail
    assert "X" * 64 not in detail
    assert detail.startswith("workbook stored column limit: sheet '\\nXXX")
    assert len(detail) < 200


_MAIN_NS = b"http://schemas.openxmlformats.org/spreadsheetml/2006/main"


async def test_xlsx_rows_are_matched_by_namespace_not_name(tmp_path):
    """TEST-SAGE-AD-244: rows are counted by the SpreadsheetML namespace the
    cell reader matches: a drawing anchor's ``<xdr:row>`` is not a stored
    row, and a worksheet row written under any prefix still is.
    """
    with_drawing = _rows_xlsx(tmp_path, [["a"]] * 5, title="Drawn")
    _add_member(
        with_drawing,
        "xl/drawings/drawing1.xml",
        b'<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/'
        b'spreadsheetDrawing"><xdr:oneCellAnchor><xdr:from><xdr:col>0</xdr:col>'
        b"<xdr:colOff>0</xdr:colOff><xdr:row>3</xdr:row><xdr:rowOff>0</xdr:rowOff>"
        b'</xdr:from><xdr:ext cx="1" cy="1"/><xdr:clientData/></xdr:oneCellAnchor>'
        b"</xdr:wsDr>",
    )
    prefixed = _rows_xlsx(tmp_path, [["a"]] * 5, title="Prefixed")
    _rewrite_member(
        prefixed,
        "xl/worksheets/sheet1.xml",
        lambda xml: xml.replace(b"<row ", b'<x:row xmlns:x="' + _MAIN_NS + b'" ').replace(
            b"</row>", b"</x:row>"
        ),
    )

    result = await XlsxAdapter().project(with_drawing, {"discovery_facts": {"max_rows": 5}})

    assert "A5 [string]\na" in result.text
    with pytest.raises(SourceReadError) as exc:
        await XlsxAdapter().project(prefixed, {"discovery_facts": {"max_rows": 4}})
    assert _guard_detail(exc) == "workbook row limit: 5 rows across the workbook; max_rows is 4"


async def test_xlsx_refusal_names_the_sheet_holding_the_row(tmp_path):
    """TEST-SAGE-AD-245: in a workbook of several sheets, a refusal names the
    sheet and part where the bound was reached, not the first sheet.
    """
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "First"
    wb.active.append(["a", "b"])
    wb.create_sheet("Second").append(["a", "b", "c"])
    path = tmp_path / "two.xlsx"
    wb.save(str(path))
    wb.close()

    with pytest.raises(SourceReadError) as exc:
        await XlsxAdapter().project(path, {"discovery_facts": {"max_columns": 2}})

    assert _guard_detail(exc) == (
        "workbook stored column limit: sheet 'Second' in part "
        "'xl/worksheets/sheet2.xml' row 1 holds 3 stored cells; max_columns is 2"
    )


_SHEET = "sheet 'Data' in part 'xl/worksheets/sheet1.xml'"


@pytest.mark.parametrize(
    ("old", "new", "limits", "expected"),
    [
        pytest.param(
            b'<c r="A1"',
            b'<c r="A123456789012345678"',
            {},
            f"workbook cell address length limit: {_SHEET} cell reference length 19; limit 16",
            id="cell-address-length",
        ),
        pytest.param(
            b'<row r="1">',
            b'<row r="1x">',
            {},
            f"workbook row address limit: {_SHEET} row reference '1x' is not a row "
            "number of at most 9 digits",
            id="row-address",
        ),
        pytest.param(
            b'<row r="1">',
            b'<row r="2000000">',
            {},
            f"workbook row limit: {_SHEET} row 2000000 exceeds the sheet maximum 1048576",
            id="row-beyond-sheet",
        ),
        pytest.param(
            b'<c r="A1"',
            b'<c r="ZZZ1"',
            {},
            f"workbook stored column limit: {_SHEET} cell 'ZZZ1' lies beyond column 16384",
            id="column-beyond-sheet",
        ),
        pytest.param(
            b"",
            b"",
            {"max_depth": 3},
            r"workbook XML depth/node limit before native allocation: "
            r"(sheet '[^']*' in )?part '[^']+' depth 4; max_depth is 3",
            id="depth",
        ),
    ],
)
async def test_xlsx_refusal_details(tmp_path, old, new, limits, expected):
    """TEST-SAGE-AD-246: every other guard refusal names where it was reached,
    the measured value and the limit it exceeded.
    """
    path = _rows_xlsx(tmp_path, [["a", "b"]])
    if old:
        _rewrite_member(path, "xl/worksheets/sheet1.xml", lambda xml: xml.replace(old, new, 1))

    with pytest.raises(SourceReadError) as exc:
        await XlsxAdapter().project(path, {"discovery_facts": limits})

    assert re.fullmatch(expected if limits else re.escape(expected), _guard_detail(exc))


async def test_xlsx_rows_and_parse_work_are_bounded_across_the_package(tmp_path):
    """TEST-SAGE-AD-247: stored rows and parse nodes are totalled across the
    whole package, so a workbook is refused when only the total exceeds the
    bound and no single sheet or part does.
    """
    import openpyxl
    from defusedxml import ElementTree as ET

    wb = openpyxl.Workbook()
    wb.active.title = "First"
    for i in range(60):
        wb.active.append([i])
    second = wb.create_sheet("Second")
    for i in range(60):
        second.append([i])
    path = tmp_path / "split.xlsx"
    wb.save(str(path))
    wb.close()
    with zipfile.ZipFile(path) as archive:
        nodes_per_part = [
            sum(1 for _ in ET.iterparse(archive.open(member)))
            for member in archive.namelist()
            if member.endswith((".xml", ".rels"))
        ]
    node_limit = max(nodes_per_part) + 1
    assert sum(nodes_per_part) > node_limit

    with pytest.raises(SourceReadError) as rows:
        await XlsxAdapter().project(path, {"discovery_facts": {"max_rows": 100}})
    with pytest.raises(SourceReadError) as nodes:
        await XlsxAdapter().project(path, {"discovery_facts": {"max_parse_nodes": node_limit}})

    assert (
        _guard_detail(rows) == "workbook row limit: 101 rows across the workbook; max_rows is 100"
    )
    assert _guard_detail(nodes) == (
        "workbook XML depth/node limit before native allocation: "
        f"{node_limit + 1} parse nodes across the package; max_parse_nodes is {node_limit}"
    )


# ---------------------------------------------------------------------------
# Office packages: decompression limits checked before parsing
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("kind", list(_PACKAGES))
async def test_package_with_an_extreme_compression_ratio_is_refused(tmp_path, kind):
    """TEST-SAGE-AD-224: a package carrying a member that inflates far beyond
    its compressed size is refused as unreadable before any parser opens it.

    The padding member is otherwise inert, so without the check every one of
    these packages projects normally.
    """
    build, adapter = _PACKAGES[kind]
    path = build(tmp_path)
    _add_member(path, "customXml/pad.bin", b"\0" * (4 * 1024 * 1024))

    with pytest.raises(SourceReadError):
        await adapter().project(path)


@pytest.mark.parametrize("kind", list(_PACKAGES))
async def test_package_over_the_total_uncompressed_limit_is_refused(tmp_path, kind, monkeypatch):
    """TEST-SAGE-AD-225: a package whose members together exceed the total
    uncompressed limit is refused, with the limit lowered so the fixture
    stays small.
    """
    monkeypatch.setattr(adapter_base, "_MAX_PACKAGE_UNCOMPRESSED_BYTES", 64 * 1024)
    build, adapter = _PACKAGES[kind]
    path = build(tmp_path)
    _add_member(path, "customXml/data.bin", bytes(range(256)) * 512)

    with pytest.raises(SourceReadError):
        await adapter().project(path)


@pytest.mark.parametrize("kind", list(_PACKAGES))
async def test_package_over_the_member_limit_is_refused(tmp_path, kind, monkeypatch):
    """TEST-SAGE-AD-236: a package with more members than the member limit
    is refused, with the limit lowered so the fixture stays small.

    The added members are tiny and compress normally, so neither the total nor
    the ratio limit can be what refuses it.
    """
    build, adapter = _PACKAGES[kind]
    path = build(tmp_path)
    with zipfile.ZipFile(path) as archive:
        present = len(archive.infolist())
    monkeypatch.setattr(adapter_base, "_MAX_PACKAGE_MEMBERS", present + 5)
    for i in range(10):
        _add_member(path, f"customXml/item{i}.xml", b"<x/>")

    with pytest.raises(SourceReadError):
        await adapter().project(path)


@pytest.mark.parametrize("kind", list(_PACKAGES))
def test_ordinary_packages_pass_the_check(tmp_path, kind):
    """TEST-SAGE-AD-226 (control): an ordinary package passes, so the limits
    do not refuse what the adapters were built for.
    """
    build, _adapter = _PACKAGES[kind]
    check_zip_package(build(tmp_path))


# ---------------------------------------------------------------------------
# Structured data: separation works within a fixed parse budget
# ---------------------------------------------------------------------------


async def test_toml_separation_is_bounded_on_header_shaped_string_lines(tmp_path):
    """TEST-SAGE-AD-227: a TOML file whose multi-line string holds thousands
    of header-shaped lines projects in bounded time and keeps the data.

    Every such line is a separation candidate the re-parse refuses, so
    halving through them re-parses the whole file roughly twice per line.
    """
    body = "\n".join(f"[h{i}]" for i in range(3000))
    source = f'title = "x"\nnotes = """\n{body}\n"""\n\n[real]\nkey = 1\n'
    path = tmp_path / "adversarial.toml"
    path.write_text(source)

    started = time.perf_counter()
    result = await StructuredDataAdapter().project(path)
    elapsed = time.perf_counter() - started

    assert elapsed < 3.0, f"{elapsed:.1f}s"
    assert "[real]" in result.text


async def test_toml_separation_still_separates_ordinary_tables(tmp_path):
    """TEST-SAGE-AD-228 (control): an ordinary TOML file still gets a blank
    line above each table, so the budget does not disable separation.
    """
    path = tmp_path / "ordinary.toml"
    path.write_text('title = "x"\n[a]\nk = 1\n[b]\nk = 2\n')

    result = await StructuredDataAdapter().project(path)

    assert "\n\n[a]\n" in result.text and "\n\n[b]\n" in result.text


# ---------------------------------------------------------------------------
# Filename parsing: linear and bounded
# ---------------------------------------------------------------------------


def test_finder_noise_strip_is_linear():
    """TEST-SAGE-AD-229: stripping Finder duplication noise from a long stem
    of repeated noise followed by a non-noise character finishes quickly.

    Searching the trailing-noise pattern over this shape tries each start and
    rescans the rest, quadratic in the stem length.
    """
    stem = " copy" * 20000 + "x"

    started = time.perf_counter()
    filename_parser._strip_finder_noise(stem)
    assert time.perf_counter() - started < 0.5


_ORIGINAL_NOISE_RE = re.compile(r"(?:[_ ]+(?:copy(?:\s+\d+)?|\(\d+\)))+\s*$")


@settings(max_examples=300, deadline=None)
@given(
    st.lists(
        st.sampled_from(
            ["a", "_", " ", "copy", " copy", "_copy 2", " (3)", "(1)", "()", "(", ")"]
            + ["2", "\u0663", "\t", "\u2003", "copy\n7", "x"]
        ),
        max_size=12,
    ).map("".join)
)
@example("a ()")
@example("a copy2")
@example("a_copy ()")
@example("a copy\u20037 (2)  ")
def test_finder_noise_strip_matches_the_pattern_it_replaces(stem):
    """TEST-SAGE-AD-230: the linear strip removes exactly what the trailing
    noise pattern matched, on stems built from the pattern's own pieces.
    """
    match = _ORIGINAL_NOISE_RE.search(stem)
    expected = stem[: match.start()] if match else stem
    assert filename_parser._strip_finder_noise(stem) == expected


def test_overlong_filename_is_not_parsed():
    """TEST-SAGE-AD-231: a stem of 1,025 characters comes back as its own
    title with nothing extracted, and quickly; one of 1,024 is parsed.
    """
    parser = FilenameParser({"filename_extraction": {"separator": "_"}})
    # A literal length rather than one read from the bound, so raising the
    # bound turns this red: 1,025 characters, one past it.
    long_stem = "2026-01-01_" + "a" * 1011 + "_v2"
    assert len(long_stem) == 1025

    started = time.perf_counter()
    parsed = parser.parse(long_stem)
    assert time.perf_counter() - started < 0.5
    assert parsed.title == long_stem and parsed.date is None and parsed.version is None

    at_bound = parser.parse("2026-01-01_" + "a" * 1010 + "_v2")
    assert at_bound.date == "2026-01-01"

    short = parser.parse("2026-01-01_Title_v2")
    assert short.date == "2026-01-01" and short.version is not None


# ---------------------------------------------------------------------------
# Projection runs off the event loop
# ---------------------------------------------------------------------------


async def test_projection_runs_on_a_worker_thread():
    """TEST-SAGE-AD-232: the ingestion service's projection helper runs the
    adapter on a ``sage-project`` worker thread, not the event-loop thread.
    """
    from sage.services.ingestion import _run_projection

    seen: list[str] = []

    class _Recording:
        async def project(self, source_path, config=None):
            seen.append(threading.current_thread().name)
            return "projected"

    assert await _run_projection(_Recording(), Path("x.md"), None) == "projected"
    assert seen and seen[0].startswith("sage-project")


def test_every_projection_in_the_ingestion_service_runs_through_the_worker():
    """TEST-SAGE-AD-235: the ingestion service never awaits an adapter's
    ``project`` directly; every projection site goes through the worker helper.

    Read from the module source, so a projection site added or reverted later
    that awaits the adapter on the event loop fails here.
    """
    import ast
    import inspect

    from sage.services import ingestion

    tree = ast.parse(inspect.getsource(ingestion))
    direct = [
        node.lineno
        for node in ast.walk(tree)
        if isinstance(node, ast.Await)
        and isinstance(node.value, ast.Call)
        and isinstance(node.value.func, ast.Attribute)
        and node.value.func.attr == "project"
    ]
    helper_calls = [
        node.lineno
        for node in ast.walk(tree)
        if isinstance(node, ast.Call)
        and isinstance(node.func, ast.Name)
        and node.func.id in {"_run_projection", "_run_inspected_projection"}
    ]
    wrapper = next(
        node
        for node in tree.body
        if isinstance(node, ast.AsyncFunctionDef) and node.name == "_run_inspected_projection"
    )
    assert any(
        isinstance(node, ast.Await)
        and isinstance(node.value, ast.Call)
        and isinstance(node.value.func, ast.Name)
        and node.value.func.id == "_run_projection"
        for node in ast.walk(wrapper)
    ), "inspection wrapper must preserve the existing projection worker boundary"
    worker = next(
        node
        for node in tree.body
        if isinstance(node, ast.AsyncFunctionDef) and node.name == "_run_projection"
    )
    materializer = next(
        node
        for node in ast.walk(worker)
        if isinstance(node, ast.FunctionDef) and node.name == "project"
    )
    guards = [
        node.lineno
        for node in ast.walk(materializer)
        if isinstance(node, ast.Call)
        and isinstance(node.func, ast.Name)
        and node.func.id == "check_projection_work"
    ]
    projections = [
        node.lineno
        for node in ast.walk(materializer)
        if isinstance(node, ast.Call)
        and isinstance(node.func, ast.Attribute)
        and node.func.attr == "project"
    ]
    assert guards and projections and max(guards) < min(projections), (
        "every original projector must pass admission on its worker before materialization"
    )
    assert direct == []
    assert len(helper_calls) >= 4


async def test_health_is_served_while_a_projection_blocks():
    """TEST-SAGE-AD-233: while an adapter blocks inside projection, the
    server answers ``/health``; the projection then completes.

    An adapter that blocks its thread on an event the test releases only
    after ``/health`` has answered: run on the event loop it stalls the loop,
    so the health request cannot be served until the block times out.
    """
    import httpx

    from sage.app import create_app
    from sage.services.ingestion import _run_projection

    started = threading.Event()
    release = threading.Event()

    class _Blocking:
        async def project(self, source_path, config=None):
            started.set()
            release.wait(timeout=2.0)
            return release.is_set()

    projection = asyncio.create_task(_run_projection(_Blocking(), Path("x.md"), None))
    while not started.is_set():
        await asyncio.sleep(0.005)

    transport = httpx.ASGITransport(app=create_app())
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
        health = await asyncio.wait_for(client.get("/health"), timeout=2.0)
    release.set()

    assert health.status_code == 200
    assert await projection is True


async def test_projection_sees_the_callers_context_variables():
    """TEST-SAGE-AD-237: the projection worker runs under a copy of the
    caller's context, so a context variable set by the request is visible to
    the adapter, and an adapter's exception reaches the caller unchanged.
    """
    import contextvars

    from sage.services.ingestion import _run_projection

    marker: contextvars.ContextVar[str] = contextvars.ContextVar("marker", default="unset")
    marker.set("request")

    class _Reading:
        async def project(self, source_path, config=None):
            return marker.get()

    class _Failing:
        async def project(self, source_path, config=None):
            raise KeyError("adapter failure")

    assert await _run_projection(_Reading(), Path("x.md"), None) == "request"
    with pytest.raises(KeyError, match="adapter failure"):
        await _run_projection(_Failing(), Path("x.md"), None)
