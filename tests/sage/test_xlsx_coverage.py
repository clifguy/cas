"""XLSX coverage must describe stored authored cells, independently of inference."""

import hashlib
import zipfile

import pytest
from openpyxl import Workbook

from sage.source_adapters.xlsx_adapter import XlsxAdapter


def workbook(tmp_path):
    path = tmp_path / "cells.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Authored > Sheet"
    ws["A1"] = "  <em>marigold</em>\n雪  "
    ws["B9"] = "beyondpreview"
    ws["XFD1048576"] = "distantcell"
    ws.row_dimensions[9].hidden = True
    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A7"] = "hiddenneedle"
    wb.save(path)
    return path


async def test_sparse_complete_hidden_and_preview_compatibility(tmp_path):
    path = workbook(tmp_path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    result = await XlsxAdapter().project(path, {"preview_rows": 1})
    assert "beyondpreview" in result.text
    assert "hiddenneedle" in result.text
    assert "distantcell" in result.text
    assert "XFD1048576" in result.text
    assert "  <em>marigold</em>\n雪  " in result.text
    assert result.content_hash == digest
    assert result.projection_complete is True
    assert result.metadata["sheet_states"]["Hidden"] == "hidden"
    assert result.metadata["hidden_rows"]["Authored > Sheet"] == ["9"]


async def test_sheet_cap_is_explicitly_limited(tmp_path):
    result = await XlsxAdapter().project(workbook(tmp_path), {"max_sheets": 1})
    assert "hiddenneedle" not in result.text
    assert result.projection_complete is False
    assert "max_sheets" in result.projection_limitations


async def test_formula_is_authored_cache_is_separate(tmp_path):
    path = tmp_path / "formula.xlsx"
    wb = Workbook()
    wb.active["A1"] = "=SUM(17,23)"
    wb.save(path)
    with zipfile.ZipFile(path) as z:
        members = {n: z.read(n) for n in z.namelist()}
    members["xl/worksheets/sheet1.xml"] = members["xl/worksheets/sheet1.xml"].replace(
        b"<v></v>", b"<v>999999</v>"
    )
    with zipfile.ZipFile(path, "w") as z:
        for name, value in members.items():
            z.writestr(name, value)
    result = await XlsxAdapter().project(path)
    assert "=SUM(17,23)" in result.text
    assert all("999999" not in p.keyword_content for p in result.passages)
    assert result.metadata["formula_caches"][0]["value"] == "999999"
    assert result.metadata["formula_caches"][0]["verified"] is False


async def test_real_storage_full_cells_and_bounded_inference(
    tmp_path,
    pg_pool,
    graph_store,
    ingestion_service_no_abstraction,
):
    from sage.adapters.content_store_postgres import PostgresContentStore
    from sage.adapters.stubs import SeededEmbeddingProvider
    from sage.models.enums import SourceType
    from sage.services.utilities import UtilitiesService
    from tests.sage.test_postgres_graph_store import _doc
    from tests.sage.test_selective_embedding import _projection_with_limits

    class Recording(SeededEmbeddingProvider):
        def __init__(self):
            super().__init__()
            self.calls = []

        async def embed(self, texts):
            self.calls.extend(texts)
            return await super().embed(texts)

    path = workbook(tmp_path)
    projection = await XlsxAdapter().project(path)
    doc = _doc(1).model_copy(
        update={"source_type": SourceType.XLSX, "adapter_version": XlsxAdapter.VERSION}
    )
    await graph_store.insert_document(doc)
    projection.discovery_evidence = _projection_with_limits(
        doc,
        [],
        max_excerpts=1,
        max_tokens=200,
        max_in_flight=1,
    ).discovery_evidence
    service = ingestion_service_no_abstraction
    service._content_store = store = PostgresContentStore(pg_pool)
    service._embedding = provider = Recording()
    await service._stage2_indexing(doc.id, projection)
    chunks = await store.get_all_chunks(doc.id)
    assert len(chunks) == 4
    assert "XFD1048576" in "".join(c.content for c in chunks)
    assert (await store.search_bm25("distantcell"))[0].document_id == doc.id
    assert (await store.search_bm25("hiddenneedle"))[0].document_id == doc.id
    assert (await store.search_bm25("marigold"))[0].document_id == doc.id
    for generated in ("XFD1048576", "string", "boolean", "unverified"):
        assert not await store.search_bm25(generated)
    reader = UtilitiesService(graph_store, store, provider, service._config)
    read = await reader.read_section(doc.id, "Authored > Sheet")
    assert "distantcell" in read.section_text and "XFD1048576" in read.section_text
    current = await graph_store.get_document(doc.id)
    assert current.processing_coverage.projection.scope == "complete"
    assert current.processing_coverage.keyword_indexing.scope == "complete"
    assert current.processing_coverage.passage_embeddings.scope == "limited"
    assert len(provider.calls) <= 1
    assert sum(provider.count_tokens(t) for t in provider.calls) <= 200
    assert len(await store.read_embedding_vectors(doc.id, list(range(4)))) <= 1
    before = list(provider.calls)
    await service._stage2_indexing(doc.id, projection)
    assert provider.calls == before, "retry must retain the cumulative allowance"
    assert len(await store.get_all_chunks(doc.id)) == 4
    assert hashlib.sha256(path.read_bytes()).hexdigest() == projection.content_hash
    from sage.services.document_surface import refresh_derived_retrieval_text

    authored_before = [c.keyword_content for c in await store.get_all_chunks(doc.id)]
    await graph_store.update_document(doc.id, {"title": "Changed source title"})
    await refresh_derived_retrieval_text(store, doc.id, await graph_store.get_document(doc.id))
    assert [c.keyword_content for c in await store.get_all_chunks(doc.id)] == authored_before
    assert not await store.search_bm25("XFD1048576")
    assert not await service._divide_stored_passages(doc.id)
    assert len(await store.get_all_chunks(doc.id)) == 4
    assert await store.migrate_indexed_structure([]) == 0
    assert not await store.search_bm25("XFD1048576")


async def test_keyword_evidence_migration_preserves_legacy_and_is_idempotent(
    pg_pool,
    graph_store,
    ingestion_service_no_abstraction,
    tmp_path,
):
    from sage.adapters.content_store_postgres import PostgresContentStore
    from sage.adapters.interfaces import Chunk
    from tests.sage.test_postgres_graph_store import _doc

    doc = _doc(1)
    await graph_store.insert_document(doc)
    store = PostgresContentStore(pg_pool)
    await store.index_chunks(doc.id, [Chunk(doc.id, "Legacy", "oldpreview")])
    # Recreate precisely the predecessor index, leaving the new nullable column
    # in place as bootstrap on an existing vault does.
    async with pg_pool.connection() as conn:
        await conn.execute("ALTER TABLE chunks DROP COLUMN tsv")
        await conn.execute(
            "ALTER TABLE chunks ADD COLUMN tsv tsvector GENERATED ALWAYS AS "
            "(to_tsvector('english', content)) STORED"
        )
    try:
        assert not await store.passage_vector_is_current()
        service = ingestion_service_no_abstraction
        service._content_store = store
        projection = await XlsxAdapter().project(workbook(tmp_path))
        await service._run_background_pipeline(doc.id, projection, doc.doc_type)
        failed = await graph_store.get_document(doc.id)
        assert failed.pipeline_status.value == "failed"
        assert "keyword_index_migration_required" in failed.pipeline_error
        assert (await store.get_all_chunks(doc.id))[0].content == "oldpreview"
        assert await store.search_bm25("oldpreview")
        assert await store.migrate_indexed_structure([]) == 0
        assert await store.passage_vector_is_current()
        legacy = (await store.get_all_chunks(doc.id))[0]
        assert legacy.keyword_content is None and legacy.content == "oldpreview"
        assert await store.search_bm25("oldpreview")
        assert await store.migrate_indexed_structure([]) == 0
        assert (await store.get_all_chunks(doc.id))[0] == legacy
        await service._stage2_indexing(doc.id, projection)
        assert await store.search_bm25("distantcell")
        assert not await store.search_bm25("XFD1048576")
    finally:
        await store.migrate_indexed_structure([])


async def test_typed_values_dates_blanks_merges_and_formula_cache_do_not_match(
    tmp_path,
    pg_pool,
    graph_store,
    ingestion_service_no_abstraction,
):
    from datetime import datetime

    from openpyxl.utils.datetime import MAC_EPOCH

    from sage.adapters.content_store_postgres import PostgresContentStore
    from tests.sage.test_postgres_graph_store import _doc

    path = tmp_path / "typed.xlsx"
    wb = Workbook()
    wb.epoch = MAC_EPOCH
    ws = wb.active
    ws["A1"] = True
    ws["B1"] = datetime(2020, 2, 3, 4, 5)
    ws["C1"] = "#DIV/0!"
    ws["D1"] = "=SUM(17,23)"
    ws["A2"] = "mergedanchor"
    ws.merge_cells("A2:C3")
    ws["D3"].number_format = "0.00"  # Explicit stored blank, absent E3 differs.
    ws.column_dimensions["B"].hidden = True
    wb.save(path)
    with zipfile.ZipFile(path) as z:
        members = {n: z.read(n) for n in z.namelist()}
    members["xl/worksheets/sheet1.xml"] = members["xl/worksheets/sheet1.xml"].replace(
        b"<v></v>", b"<v>stalephantom</v>"
    )
    with zipfile.ZipFile(path, "w") as z:
        for name, value in members.items():
            z.writestr(name, value)
    projection = await XlsxAdapter().project(path)
    assert "2020-02-03T04:05:00" in projection.text
    assert "[boolean]\ntrue" in projection.text
    assert "[error]\n#DIV/0!" in projection.text
    assert "[blank]" in projection.text
    assert "E3" not in projection.text
    assert projection.text.count("mergedanchor") == 1
    assert projection.metadata["merges"]["Sheet"] == ["A2:C3"]
    assert projection.metadata["epoch"].startswith("1904-01-01")
    assert projection.metadata["hidden_columns"]["Sheet"] == [("2", "2")]
    doc = _doc(1)
    await graph_store.insert_document(doc)
    service = ingestion_service_no_abstraction
    service._content_store = store = PostgresContentStore(pg_pool)
    await service._stage2_indexing(doc.id, projection)
    assert await store.search_bm25("SUM")
    for term in ("stalephantom", "cache", "boolean", "blank"):
        assert not await store.search_bm25(term)
    assert "stalephantom" in "".join(c.content for c in await store.get_all_chunks(doc.id))


async def test_large_numeric_workbook_ingest_recovery_and_wire_parity(
    tmp_path,
    minimal_config,
    pg_pool,
    graph_store,
):
    import json
    from urllib.parse import quote

    from httpx import ASGITransport, AsyncClient

    from sage.adapters.content_store_postgres import PostgresContentStore
    from sage.adapters.stubs import SeededEmbeddingProvider, StubAbstractionProvider
    from sage.app import create_app
    from sage.mcp_server import _vaults, build_partitioned_server
    from sage.models.schemas import IngestRequest
    from tests.helpers.pipeline_wait import drain_abstraction_queue
    from tests.sage.conftest import initialize_services_for_test

    class Embedding(SeededEmbeddingProvider):
        def __init__(self):
            super().__init__()
            self.inputs = []

        async def embed(self, texts):
            self.inputs.extend(texts)
            return await super().embed(texts)

    class Abstraction(StubAbstractionProvider):
        def __init__(self):
            self.calls = []

        async def generate_abstract(self, text, max_tokens, doc_type):
            self.calls.append((str(text), max_tokens))
            return "Retained authored workbook evidence."

    source = tmp_path / "numeric.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Cells > precise"
    for row in range(1, 1001):
        ws.append([row, "repeated sentence about saffron.", f"near duplicate sentence {row % 7}"])
    ws["C1001"] = "lateauthoredneedle"
    wb.save(source)
    config = minimal_config.model_copy(deep=True)
    config.abstraction.max_generation_calls = 1
    provider, abstraction = Embedding(), Abstraction()
    store = PostgresContentStore(pg_pool)
    async with initialize_services_for_test(
        config,
        graph_store=graph_store,
        content_store=store,
        embedding_provider=provider,
        abstraction_provider=abstraction,
    ) as services:
        service = services.ingestion_service
        request = IngestRequest(
            source=str(source),
            source_type="xlsx",
            config={
                "discovery_selection": {"max_excerpts": 2, "max_tokens": 256, "max_in_flight": 1},
            },
        )
        result = await service.ingest(request)
        doc = result.document
        assert doc.processing_coverage.projection.scope == "complete"
        assert doc.processing_coverage.keyword_indexing.scope == "complete"
        assert doc.processing_coverage.passage_embeddings.scope == "limited"
        assert len(await store.get_all_chunks(doc.id)) == 3001
        assert len(provider.inputs) <= 2
        assert sum(provider.count_tokens(t) for t in provider.inputs) <= 256
        assert len(abstraction.calls) == 1
        selected_text, reserved_output = abstraction.calls[0]
        assert len(selected_text.encode()) <= 256
        assert reserved_output <= config.abstraction.max_total_output_tokens
        from sage.adapters.abstraction_prompt import _format_system_prompt, wrap_source_document

        protected_input = (
            len(
                (
                    _format_system_prompt(doc.doc_type, bounded=True)
                    + wrap_source_document(selected_text)
                ).encode()
            )
            + 1024
        )
        output_upper_bound = len("Retained authored workbook evidence.".encode())
        assert protected_input <= config.abstraction.max_total_input_tokens
        assert output_upper_bound <= reserved_output
        assert (
            "lateauthoredneedle"
            in (await services.utilities_service.read_projection(doc.id)).projection_text
        )
        # Both thin transports read/search the same durable generation. The app
        # consumes these Core API operations, with no workbook preview renderer.
        ordinary = build_partitioned_server("sage")
        _vaults[config.vault.id] = services
        app = create_app(config=config)
        app.state.vault_registry = {config.vault.id: services}
        try:
            async with AsyncClient(
                transport=ASGITransport(app=app), base_url="http://test"
            ) as client:
                base = f"/sage_vaults/{config.vault.id}"
                rest = await client.get(f"{base}/documents/{doc.id}/projection")
                assert rest.status_code == 200
                mcp = json.loads(
                    (
                        await ordinary.call_tool(
                            "read_projection",
                            {
                                "vault_id": config.vault.id,
                                "document_id": doc.id,
                            },
                        )
                    )[0].text
                )
                assert mcp["projection_text"] == rest.json()["projection_text"]
                heading = quote(ws.title, safe="")
                section = await client.get(f"{base}/documents/{doc.id}/section/{heading}")
                assert (
                    section.status_code == 200
                    and "lateauthoredneedle" in section.json()["section_text"]
                )
                mcp_search = json.loads(
                    (
                        await ordinary.call_tool(
                            "search",
                            {
                                "vault_id": config.vault.id,
                                "mode": "keyword",
                                "query": "lateauthoredneedle",
                            },
                        )
                    )[0].text
                )
                assert (
                    mcp_search["results"] and mcp_search["results"][0]["document"]["id"] == doc.id
                )
                for term in ("unverified", "string", "C1001"):
                    negative = json.loads(
                        (
                            await ordinary.call_tool(
                                "search",
                                {
                                    "vault_id": config.vault.id,
                                    "mode": "keyword",
                                    "query": term,
                                },
                            )
                        )[0].text
                    )
                    assert not negative["results"]
        finally:
            _vaults.pop(config.vault.id, None)
        before = list(provider.inputs)
        await service.recompute_pipeline(doc.id)
        await drain_abstraction_queue(service)
        current = await services.graph_store.get_document(doc.id)
        assert current.processing_coverage.projection.scope == "complete"
        assert provider.inputs == before
        assert len(abstraction.calls) == 1
        await service.reabstract(doc.id)
        await drain_abstraction_queue(service)
        assert len(abstraction.calls) == 1
        assert len(await store.get_all_chunks(doc.id)) == 3001
        assert await store.search_bm25("lateauthoredneedle")
        # Emit actual deterministic provider measurements alongside pytest output.
        print(
            json.dumps(
                {
                    "cells": 3001,
                    "embedding_calls_inputs": len(provider.inputs),
                    "embedding_tokens": sum(provider.count_tokens(t) for t in provider.inputs),
                    "abstraction_calls": len(abstraction.calls),
                    "abstraction_selected_input_bytes": len(selected_text.encode()),
                    "abstraction_reserved_output_tokens": reserved_output,
                    "abstraction_protected_input_token_upper_bound": protected_input,
                    "abstraction_output_token_upper_bound": output_upper_bound,
                }
            )
        )


@pytest.mark.parametrize("keyword_cap,fail", [(15, False), (None, True)])
async def test_workbook_independent_limits_partial_failure_and_stale_fence(
    tmp_path,
    pg_pool,
    graph_store,
    ingestion_service_no_abstraction,
    keyword_cap,
    fail,
):
    from sage.adapters.content_store_postgres import PostgresContentStore
    from sage.adapters.stubs import SeededEmbeddingProvider
    from tests.sage.test_postgres_graph_store import _doc
    from tests.sage.test_selective_embedding import _projection_with_limits

    class Controlled(SeededEmbeddingProvider):
        def __init__(self):
            super().__init__()
            self.calls = 0

        async def embed(self, texts):
            self.calls += 1
            if fail and self.calls == 2:
                raise RuntimeError("controlled workbook provider failure")
            return await super().embed(texts)

    path = tmp_path / "limits.xlsx"
    wb = Workbook()
    for row in range(1, 6):
        wb.active.cell(row, 1, f"authoredneedle{row}")
    wb.save(path)
    projection = await XlsxAdapter().project(path)
    doc = _doc(1)
    await graph_store.insert_document(doc)
    projection.discovery_evidence = _projection_with_limits(
        doc,
        [],
        max_excerpts=5,
        max_tokens=500,
        max_in_flight=1,
        max_keyword_characters=keyword_cap,
    ).discovery_evidence
    service = ingestion_service_no_abstraction
    service._content_store = store = PostgresContentStore(pg_pool)
    service._embedding = Controlled()
    if fail:
        with pytest.raises(RuntimeError, match="controlled workbook"):
            await service._stage2_indexing(doc.id, projection)
    else:
        await service._stage2_indexing(doc.id, projection)
    coverage, chunks = await store.read_processing_snapshot(doc.id)
    assert coverage.projection.scope == "complete" and len(chunks) == 5
    assert coverage.keyword_indexing.scope == ("complete" if fail else "limited")
    assert bool(await store.search_bm25("authoredneedle5")) == fail
    if fail:
        assert 0 < len(await store.read_embedding_vectors(doc.id, list(range(5)))) < 5
    old = coverage.attempt
    newer = await store.begin_processing_attempt(doc.id)
    assert newer > old
    from dataclasses import replace

    from sage.models.coverage import CoverageDimension

    stale_chunks = [replace(c, content="stale producer overwrite") for c in chunks]
    assert not await store.publish_processing(
        doc.id,
        old,
        {
            name: CoverageDimension()
            for name in ("projection", "keyword_indexing", "passage_embeddings")
        },
        chunks=stale_chunks,
    )
    served = await store.get_all_chunks(doc.id)
    assert len(served) == 5 and all(c.content != "stale producer overwrite" for c in served)


async def test_sparse_dimensions_parser_limits_and_lexical_number(tmp_path):
    from sage.source_adapters.base import SourceReadError

    path = workbook(tmp_path)
    with zipfile.ZipFile(path) as z:
        members = {n: z.read(n) for n in z.namelist()}
    members["xl/worksheets/sheet1.xml"] = (
        members["xl/worksheets/sheet1.xml"]
        .replace(
            b'<dimension ref="A1:XFD1048576"/>',
            b'<dimension ref="A1:A1"/>',
        )
        .replace(
            b"</sheetData>",
            b'<row r="42"><c r="B42" t="n"><v>'
            b"1.2345678901234567890123456789</v></c></row></sheetData>",
        )
    )
    with zipfile.ZipFile(path, "w") as z:
        for name, value in members.items():
            z.writestr(name, value)
    result = await XlsxAdapter().project(path)
    assert "distantcell" in result.text
    assert "1.2345678901234567890123456789" in result.text
    assert len(result.passages) == 5
    with pytest.raises(SourceReadError, match="row limit"):
        await XlsxAdapter().project(path, {"discovery_facts": {"max_rows": 1}})


async def test_csv_and_xlsx_authored_content_are_reachable(
    tmp_path, pg_pool, graph_store, ingestion_service_no_abstraction
):
    from sage.adapters.content_store_postgres import PostgresContentStore
    from sage.source_adapters.structured_data_adapter import StructuredDataAdapter
    from tests.sage.test_postgres_graph_store import _doc

    path = tmp_path / "equivalent.xlsx"
    wb = Workbook()
    wb.active.append(["<markup>cedar</markup>", "雪"])
    wb.active["A9"] = "lateequivalent"
    wb.save(path)
    csv = tmp_path / "equivalent.csv"
    csv.write_text("<markup>cedar</markup>,雪\n" + "\n" * 7 + "lateequivalent,\n")
    store = PostgresContentStore(pg_pool)
    service = ingestion_service_no_abstraction
    service._content_store = store
    for index, projection in enumerate(
        (await XlsxAdapter().project(path), await StructuredDataAdapter().project(csv)), 1
    ):
        doc = _doc(index)
        await graph_store.insert_document(doc)
        await service._stage2_indexing(doc.id, projection)
    for term in ("cedar", "lateequivalent"):
        assert {r.document_id for r in await store.search_bm25(term)} == {_doc(1).id, _doc(2).id}
    assert not await store.search_bm25("A9")


async def test_historical_preview_is_not_automatically_reprojected(
    graph_store,
    pg_pool,
    ingestion_service_no_abstraction,
):
    from sage.adapters.content_store_postgres import PostgresContentStore
    from sage.adapters.interfaces import Chunk
    from sage.models.enums import SourceType
    from tests.sage.test_postgres_graph_store import _doc

    class Forbidden(XlsxAdapter):
        async def project(self, *args, **kwargs):
            pytest.fail("historical source reprocessing was not requested")

    doc = _doc(1).model_copy(update={"source_type": SourceType.XLSX, "adapter_version": "0.1.0"})
    await graph_store.insert_document(doc)
    store = PostgresContentStore(pg_pool)
    await store.index_chunks(doc.id, [Chunk(doc.id, " > legacy", "historicpreview")])
    before = (await graph_store.get_document(doc.id)).processing_coverage
    service = ingestion_service_no_abstraction
    service._content_store = store
    service._adapters[SourceType.XLSX] = Forbidden()
    assert await service.store_text_before_first_heading() == (0, [])
    after = await graph_store.get_document(doc.id)
    assert after.adapter_version == "0.1.0" and after.processing_coverage == before
    assert (await store.get_all_chunks(doc.id))[0].content == "historicpreview"


async def test_workbook_cancellation_preserves_readable_generation_and_allowance(
    tmp_path,
    pg_pool,
    graph_store,
    ingestion_service_no_abstraction,
):
    import asyncio

    from sage.adapters.content_store_postgres import PostgresContentStore
    from sage.adapters.stubs import SeededEmbeddingProvider
    from tests.sage.test_postgres_graph_store import _doc
    from tests.sage.test_selective_embedding import _projection_with_limits

    class Blocking(SeededEmbeddingProvider):
        def __init__(self):
            super().__init__()
            self.entered = asyncio.Event()
            self.calls = 0

        async def embed(self, texts):
            self.calls += 1
            self.entered.set()
            await asyncio.Future()

    projection = await XlsxAdapter().project(workbook(tmp_path))
    doc = _doc(1)
    await graph_store.insert_document(doc)
    projection.discovery_evidence = _projection_with_limits(
        doc, [], max_excerpts=1
    ).discovery_evidence
    service = ingestion_service_no_abstraction
    service._content_store = store = PostgresContentStore(pg_pool)
    service._embedding = provider = Blocking()
    task = asyncio.create_task(service._stage2_indexing(doc.id, projection))
    await asyncio.wait_for(provider.entered.wait(), 5)
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert await store.search_bm25("distantcell")
    assert len(await store.get_all_chunks(doc.id)) == 4
    await asyncio.wait_for(service._stage2_indexing(doc.id, projection), 5)
    assert provider.calls == 1
    assert (
        await graph_store.get_document(doc.id)
    ).processing_coverage.projection.scope == "complete"


def shared_string_workbook(tmp_path, references):
    path = tmp_path / f"shared-{references}.xlsx"
    book = Workbook()
    book.save(path)
    with zipfile.ZipFile(path) as z:
        members = {name: z.read(name) for name in z.namelist()}
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rows = "".join(
        f'<row r="{i}"><c r="A{i}" t="s"><v>0</v></c></row>' for i in range(1, references + 1)
    )
    members["xl/worksheets/sheet1.xml"] = (
        f'<worksheet xmlns="{namespace}"><sheetData>{rows}</sheetData></worksheet>'.encode()
    )
    members["xl/sharedStrings.xml"] = (
        f'<sst xmlns="{namespace}"><si><t>{"s" * 32760}</t></si></sst>'.encode()
    )
    members["[Content_Types].xml"] = members["[Content_Types].xml"].replace(
        b"</Types>",
        b'<Override PartName="/xl/sharedStrings.xml" '
        b'ContentType="application/vnd.openxmlformats-officedocument.'
        b'spreadsheetml.sharedStrings+xml"/></Types>',
    )
    with zipfile.ZipFile(path, "w") as z:
        for name, value in members.items():
            z.writestr(name, value)
    return path


async def test_shared_string_expansion_is_bounded_before_display_allocation(tmp_path):
    from sage.source_adapters.base import SourceReadError
    from sage.source_adapters.structure import FactLimits, check_projection_work

    small = shared_string_workbook(tmp_path, 2)
    result = await XlsxAdapter().project(small)
    assert len(result.passages) == 2
    assert sum(len(p.keyword_content) for p in result.passages) == 2 * 32760
    expanded = shared_string_workbook(tmp_path, 600)
    assert expanded.stat().st_size < 100_000
    check_projection_work(expanded, FactLimits())  # Source/XML bounds admit it.
    with pytest.raises(SourceReadError, match="projection character limit"):
        await XlsxAdapter().project(expanded)


async def test_array_and_data_table_formula_attributes_are_readable_not_authored(tmp_path):
    path = tmp_path / "formula-attributes.xlsx"
    wb = Workbook()
    wb.active["A1"] = "=SUM(17,23)"
    wb.active["B1"] = "=SUM(1,2)"
    wb.save(path)
    with zipfile.ZipFile(path) as z:
        members = {n: z.read(n) for n in z.namelist()}
    members["xl/worksheets/sheet1.xml"] = (
        members["xl/worksheets/sheet1.xml"]
        .replace(
            b"<f>SUM(17,23)</f>",
            b'<f t="array" ref="A1:A3">SUM(17,23)</f>',
        )
        .replace(b"<f>SUM(1,2)</f>", b'<f t="dataTable" ref="B1:B3" dt2D="1" r1="D1"/>')
    )
    with zipfile.ZipFile(path, "w") as z:
        for name, value in members.items():
            z.writestr(name, value)
    result = await XlsxAdapter().project(path)
    assert "A1:A3" in result.text and "B1:B3" in result.text and "dt2D" in result.text
    assert "=SUM(17,23)" in result.passages[0].keyword_content
    assert result.passages[1].keyword_content == ""
    assert all(
        "dataTable" not in p.keyword_content and "dt2D" not in p.keyword_content
        for p in result.passages
    )
