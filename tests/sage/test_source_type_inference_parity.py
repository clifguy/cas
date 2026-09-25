"""An omitted ``source_type`` is inferred beneath every request surface.

Inference from the source's extension is a capability of the ingest operation,
not of the MCP tool that first offered it (CAS-ADR-052), so it runs inside the
ingestion service where the Core API's JSON ingest, the MCP tool, and both batch
surfaces reach it on the same terms. Precedence is explicit, then extension: a
value the caller supplies is never replaced, and an extension no registered
adapter claims is refused as ``source_type_unresolved`` rather than guessed.

The inference reads the path the delivery resolved to. On a ``transfer_token``
completion that is the staged file, which carries the caller's own basename, so
a completion naming no type infers exactly as the co-located call would have.
"""

import contextlib
import io
import json
from pathlib import Path

import docx
import pytest
import yaml
from httpx import ASGITransport, AsyncClient

import sage.mcp_init as _mcp_init
import sage.mcp_server as _mcp
from sage.adapters.stubs import (
    StubAbstractionProvider,
    StubContentStore,
    StubEmbeddingProvider,
)
from sage.app import _initialize_services, create_app
from sage.config import SageCoreConfig, VaultConfig
from sage.mcp_server import bulk_ingest_document, ingest_document, search
from sage.models.schemas import BatchIngestFileMetadata, IngestRequest
from sage.services.transfer import get_transfer_store, reset_transfer_store
from tests.helpers.pipeline_wait import drain_vaults

_VAULT_ID = "test_vault"
_BASE = "https://sage.test.example"
_INGEST = f"/sage_vaults/{_VAULT_ID}/documents"
_BATCH = f"/sage_vaults/{_VAULT_ID}/documents:batch"
_SPEC = Path(__file__).resolve().parents[2] / "docs" / "fs" / "sage" / "sage_core_api.openapi.yaml"


@contextlib.contextmanager
def _profile(name: str):
    """Pin the deployment profile and transfer coordinates for the block."""
    saved = _mcp_init._stack_config
    _mcp_init.set_stack_config(SageCoreConfig(profile=name, transfer={"public_base_url": _BASE}))
    try:
        yield
    finally:
        _mcp_init.set_stack_config(saved)


def _parse(result: str | dict) -> dict:
    return result if isinstance(result, dict) else json.loads(result)


def _summary_of(text: str) -> dict:
    events = [
        json.loads(line.replace("data: ", "", 1))
        for line in text.strip().split("\n")
        if line.startswith("data: ")
    ]
    return next(e for e in events if e["event_type"] == "summary")


def _docx_bytes() -> bytes:
    document = docx.Document()
    document.add_heading("Inferred", level=1)
    document.add_paragraph("Body.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture(autouse=True)
def _fresh_transfer_store():
    reset_transfer_store()
    yield
    reset_transfer_store()


@pytest.fixture
async def vault(minimal_vault_config_dict):
    """One initialized vault, shared between the HTTP app and the MCP tools."""
    config = VaultConfig.model_validate(minimal_vault_config_dict)
    application = create_app(config=config)
    await _initialize_services(
        application,
        config,
        content_store=StubContentStore(),
        embedding_provider=StubEmbeddingProvider(),
        abstraction_provider=StubAbstractionProvider(),
    )
    try:
        yield application, Path(config.vault.storage_root)
    finally:
        try:
            await drain_vaults(application.state.vault_registry)
        finally:
            for services in application.state.vault_registry.values():
                services.close_timing()
                await services.close_storage()
            _mcp._vaults.pop(_VAULT_ID, None)


@pytest.fixture
async def client(vault):
    application, _root = vault
    async with AsyncClient(transport=ASGITransport(app=application), base_url="http://test") as c:
        yield c


def _registered_types(vault) -> list[str]:
    application, _root = vault
    service = application.state.vault_registry[_VAULT_ID].ingestion_service
    return sorted(source_type.value for source_type in service._adapters)  # noqa: SLF001


def _registered_extensions(vault) -> dict[str, list[str]]:
    application, _root = vault
    service = application.state.vault_registry[_VAULT_ID].ingestion_service
    return {
        source_type.value: list(adapter.EXTENSIONS)
        for source_type, adapter in service._adapters.items()  # noqa: SLF001
    }


def _write(root: Path, relative: str, body: bytes) -> str:
    path = root / relative
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(body)
    return relative


def _caller_file(tmp_path: Path, name: str, body: bytes) -> Path:
    src = tmp_path / "caller_inbox" / name
    src.parent.mkdir(parents=True, exist_ok=True)
    src.write_bytes(body)
    return src


async def _staged_token(client: AsyncClient, src: Path, body: bytes) -> tuple[str, Path]:
    """Mint a recipe naming no type, deliver the bytes, and return the token.

    The completion sends the token alone. The staged file is named after the
    caller's own basename, so an inference read from either spelling finds the
    same extension; the completion cannot tell them apart and does not try to.
    """
    minted = await client.post(_INGEST, json={"source": str(src)})
    assert minted.status_code == 200, minted.text
    (item,) = minted.json()["uploads"]
    staging_dir = get_transfer_store()._entries[item["transfer_id"]].staging_dir
    resp = await client.put("/upload", content=body, headers={"X-Upload-Token": item["token"]})
    assert resp.status_code == 201, resp.text
    return item["token"], staging_dir


async def _document_count() -> int:
    return _parse(await search(_VAULT_ID, mode="catalog", limit=0))["total_available"]


# ---------------------------------------------------------------------------
# Inference on the Core API's JSON ingest
# ---------------------------------------------------------------------------


async def test_rest_ingest_without_source_type_infers_from_extension(vault, client):
    """A ``.md`` source ingests with no ``source_type`` on the request."""
    _application, root = vault
    source = _write(root, "test/note.md", b"# Note\n\nBody.\n")

    resp = await client.post(_INGEST, json={"source": source})

    assert resp.status_code == 201, resp.text
    assert resp.json()["document"]["source_type"] == "markdown"


async def test_rest_ingest_infers_a_non_markdown_type(vault, client):
    """A ``.docx`` source resolves to docx, so inference is not pinned to markdown."""
    _application, root = vault
    source = _write(root, "test/report.docx", _docx_bytes())

    resp = await client.post(_INGEST, json={"source": source})

    assert resp.status_code == 201, resp.text
    assert resp.json()["document"]["source_type"] == "docx"


async def test_xml_source_infers_structured_data_on_both_surfaces(vault, client):
    """A ``.xml`` source with no ``source_type`` resolves to structured data on both surfaces."""
    _application, root = vault
    body = '<scan surface="{}"><host addr="10.0.0.1"/>\n<host addr="10.0.0.2"/></scan>\n'
    via_mcp = _write(root, "test/scan_mcp.xml", body.format("mcp").encode())
    via_rest = _write(root, "test/scan_rest.xml", body.format("rest").encode())

    mcp = _parse(await ingest_document(_VAULT_ID, source=via_mcp))
    rest = await client.post(_INGEST, json={"source": via_rest})

    assert "error" not in mcp, mcp
    assert mcp["source_type"] == "structured_data", mcp
    assert rest.status_code == 201, rest.text
    assert rest.json()["document"]["source_type"] == "structured_data"


async def test_rest_transfer_token_completion_infers_from_staged_basename(client, tmp_path):
    """A completion carrying only the token has a type inferred at all.

    The staged file carries the caller's basename, so this shows inference runs on
    the completion leg; it does not distinguish the staged spelling from the
    declared one, which share an extension by construction.
    """
    body = b"# Staged\n\nDelivered over the upload leg.\n"
    src = _caller_file(tmp_path, "staged_note.md", body)

    with _profile("cloud"):
        token, _staging_dir = await _staged_token(client, src, body)
        done = await client.post(_INGEST, json={"transfer_token": token})

    assert done.status_code == 201, done.text
    document = done.json()["document"]
    assert document["source_type"] == "markdown"
    assert document["source_path"] == "imports/staged_note.md"


async def test_rest_explicit_type_disagreeing_with_extension_is_honored(vault, client):
    """An explicit ``source_type`` is never replaced by what the extension says.

    The ``.txt`` arm succeeds only if the explicit value is honored, since no
    adapter claims the extension. It cannot tell "explicit, else inferred" from
    "inferred, else explicit", because inference yields nothing there. The
    arm of a markdown body named ``.pdf`` and declared ``markdown`` does: under
    the inverted precedence it would be routed to the pdf adapter.
    """
    _application, root = vault
    txt = _write(root, "test/plain.txt", b"# Plain\n\nBody.\n")
    pdf_named = _write(root, "test/declared.pdf", b"# Declared\n\nBody.\n")

    honored = await client.post(_INGEST, json={"source": txt, "source_type": "markdown"})
    disagreeing = await client.post(_INGEST, json={"source": pdf_named, "source_type": "markdown"})

    assert honored.status_code == 201, honored.text
    assert honored.json()["document"]["source_type"] == "markdown"
    assert disagreeing.status_code == 201, disagreeing.text
    assert disagreeing.json()["document"]["source_type"] == "markdown"


# ---------------------------------------------------------------------------
# The unresolvable type, refused once for both surfaces
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("delivery", "name", "extension"),
    [
        ("relative", "unclaimed.txt", ".txt"),
        ("relative", "README", None),
        ("transfer", "blob.xyz", ".xyz"),
    ],
    ids=["unclaimed-extension", "no-extension", "transfer-completion"],
)
async def test_unresolvable_type_is_refused_identically_on_both_surfaces(
    vault, client, tmp_path, delivery, name, extension
):
    """The same unresolvable ingest is refused with one code, message and detail.

    Compared surface against surface and also against the literal code, so two
    surfaces agreeing on the wrong envelope fail too. Nothing is created.
    """
    _application, root = vault
    body = b"opaque bytes\n"
    before = await _document_count()

    if delivery == "relative":
        source = _write(root, f"test/{name}", body)
        mcp = _parse(await ingest_document(_VAULT_ID, source=source))
        rest = await client.post(_INGEST, json={"source": source})
    else:
        with _profile("cloud"):
            mcp_token, _ = await _staged_token(client, _caller_file(tmp_path, name, body), body)
            rest_token, _ = await _staged_token(client, _caller_file(tmp_path, name, body), body)
            mcp = _parse(await ingest_document(_VAULT_ID, transfer_token=mcp_token))
            rest = await client.post(_INGEST, json={"transfer_token": rest_token})

    assert rest.status_code == 400, rest.text
    envelope = rest.json()
    assert envelope["code"] == "source_type_unresolved", envelope
    assert envelope["detail"] == {
        "extension": extension,
        "registered_source_types": _registered_types(vault),
        "registered_extensions": _registered_extensions(vault),
    }
    assert mcp.get("error") == envelope["code"], mcp
    assert mcp["message"] == envelope["message"]
    assert mcp["detail"] == envelope["detail"]
    assert await _document_count() == before


async def test_unresolvable_type_refusal_does_not_disclose_staging_path(client, tmp_path):
    """The refusal names the extension, never the location the bytes were staged at."""
    body = b"opaque bytes\n"
    src = _caller_file(tmp_path, "hidden.xyz", body)

    with _profile("cloud"):
        token, staging_dir = await _staged_token(client, src, body)
        rest = await client.post(_INGEST, json={"transfer_token": token})

    assert rest.status_code == 400, rest.text
    rendered = json.dumps(rest.json())
    assert rest.json()["code"] == "source_type_unresolved"
    assert str(staging_dir) not in rendered
    assert "caller_inbox" not in rendered


async def test_dry_run_reports_the_unresolved_refusal(vault, client):
    """A preview raises the refusal the real ingest would, rather than previewing it."""
    _application, root = vault
    source = _write(root, "test/preview.txt", b"# Preview\n")

    resp = await client.post(_INGEST, json={"source": source, "dry_run": True})

    assert resp.status_code == 400, resp.text
    assert resp.json()["code"] == "source_type_unresolved"


# ---------------------------------------------------------------------------
# Both batch surfaces
# ---------------------------------------------------------------------------


async def test_bulk_mcp_entry_without_source_type_infers_and_refuses_per_file(vault):
    """Entries naming no type infer per file, and an unclaimed one is refused alone."""
    _application, root = vault
    md = _write(root, "test/bulk_a.md", b"# A\n\nBody.\n")
    txt = _write(root, "test/bulk_b.txt", b"opaque\n")

    summary = _parse(
        await bulk_ingest_document(
            _VAULT_ID, [{"file_path": md}, {"file_path": txt}], infer_edges=False
        )
    )

    assert summary.get("error") is None, summary
    assert summary["error_count"] == 1, summary
    (entry,) = summary["errors"]
    assert entry["file_index"] == 1, entry
    assert entry["code"] == "source_type_unresolved", entry
    assert entry["source_path"] == txt, entry
    created = _parse(
        await search(_VAULT_ID, mode="catalog", filters={"source_type": "markdown"}, limit=10)
    )
    assert created["total_available"] == 1, created


async def test_multipart_batch_entry_without_source_type_infers_and_refuses_per_file(client):
    """The multipart batch applies the same rule to descriptors naming no type."""
    metadata = json.dumps(
        {
            "infer_edges": False,
            "files": [{"parsed_metadata": {"title": "A"}}, {"parsed_metadata": {"title": "B"}}],
        }
    )

    resp = await client.post(
        _BATCH,
        files=[
            ("files", ("upload_a.md", b"# A\n\nBody.\n", "text/markdown")),
            ("files", ("upload_b.txt", b"opaque\n", "text/plain")),
        ],
        data={"metadata": metadata},
    )

    assert resp.status_code == 200, resp.text
    summary = _summary_of(resp.text)
    assert summary["documents_created"]["new"] == 1, summary
    (error,) = summary["errors"]
    assert error["file_index"] == 1, error
    assert error.get("code") == "source_type_unresolved", error


# ---------------------------------------------------------------------------
# The request contract
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("schema_name", "model"),
    [("IngestRequest", IngestRequest), ("BatchIngestFileMetadata", BatchIngestFileMetadata)],
)
def test_source_type_is_optional_in_spec_and_model(schema_name, model):
    """The formal substrate and the model derived from it both admit an omitted type.

    Both also admit an explicit null. The model's ``| None`` accepts one, so a
    spec property with no null arm states a narrower contract than the surface
    serves -- which the ``required`` list alone cannot show.
    """
    schema = yaml.safe_load(_SPEC.read_text())["components"]["schemas"][schema_name]
    prop = schema["properties"]["source_type"]
    types = prop.get("type")
    admits_null = (isinstance(types, list) and "null" in types) or any(
        arm.get("type") == "null" for arm in prop.get("anyOf", [])
    )

    assert "source_type" not in schema.get("required", []), schema.get("required")
    assert admits_null, prop
    field = model.model_fields["source_type"]
    assert not field.is_required()
    assert field.default is None
    body = {"source_type": None} | ({"source": "a.md"} if model is IngestRequest else {})
    assert model.model_validate(body).source_type is None
