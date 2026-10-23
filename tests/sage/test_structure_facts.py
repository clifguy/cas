"""Behavioral tests for the shared native-facts seam."""

from pathlib import Path

import pytest

from sage.source_adapters.base import HeadingNode, ProjectionResult
from sage.source_adapters.structure import FactLimits, source_facts


def facts(tmp_path: Path, suffix: str, text: str, **limits):
    path = tmp_path / f"source{suffix}"
    path.write_text(text)
    stream = source_facts(path, None, FactLimits(**limits))
    return list(stream), stream


@pytest.mark.parametrize(
    ("suffix", "text", "expected"),
    [
        (".json", '{"value": [true, 2, "{\\"a\\":1}"]}', [True, 2, '{"a":1}']),
        (".jsonl", '{"v":1}\n{"v":"late text"}\n', [1, "late text"]),
        (".yaml", 'value: [true, 2, "hello"]', [True, 2, "hello"]),
        (".yml", 'value: [true, 2, "hello"]', [True, 2, "hello"]),
        (".toml", 'value = [true, 2, "hello"]', [True, 2, "hello"]),
        (".csv", 'name,value\n"multi\nline",2\n', ["name", "value", "multi\nline", "2"]),
        (".tsv", "name\tvalue\na\t2\n", ["name", "value", "a", "2"]),
    ],
)
def test_native_types_and_lexical_strings(tmp_path, suffix, text, expected):
    actual, stream = facts(tmp_path, suffix, text)
    assert [fact.value for fact in actual] == expected
    assert stream.exhausted
    assert not stream.limitations or any(
        "embedded structure unresolved" in item for item in stream.limitations
    )


def test_csv_uses_ordinals_and_physical_line_locations(tmp_path):
    actual, _ = facts(tmp_path, ".csv", 'title,value\n"a\nb",2\nc,3\n')
    assert [fact.field_path for fact in actual] == ["column:1", "column:2"] * 3
    assert actual[-2].location == "line:4:column:1"
    assert actual[2].record == "row:2"


def test_json_guard_precedes_native_tree_allocation(tmp_path, monkeypatch):
    import json

    calls = []
    monkeypatch.setattr(json, "loads", lambda source: calls.append(source))
    actual, stream = facts(tmp_path, ".json", "[[[[1]]]]", max_depth=2)
    assert not calls
    assert not actual
    assert "before tree allocation" in stream.limitations[0]
    assert not stream.exhausted


def test_yaml_aliases_are_unexpanded_authored_source(tmp_path):
    text = "a: &a [hello]\nb: &b [*a, *a]\nc: [*b, *b]\n"
    actual, stream = facts(tmp_path, ".yaml", text)
    assert [fact.rendering for fact in actual] == text.splitlines()
    assert "expanded native values uninspected" in stream.limitations[0]


def test_row_cutoff_does_not_read_remainder(tmp_path):
    actual, stream = facts(tmp_path, ".csv", 'first\nsecond\n"invalid', max_rows=2)
    assert [fact.value for fact in actual] == ["first", "second"]
    assert "remainder unknown" in stream.limitations[0]
    assert not stream.exhausted


def test_byte_admission_precedes_read(tmp_path, monkeypatch):
    path = tmp_path / "source.json"
    path.write_text('["hello"]')
    monkeypatch.setattr(Path, "read_text", lambda *a, **k: pytest.fail("source was read"))
    stream = source_facts(path, None, FactLimits(max_source_bytes=2))
    assert list(stream) == []
    assert "not inspected" in stream.limitations[0]


def test_xml_mixed_content_and_attributes(tmp_path):
    actual, stream = facts(tmp_path, ".xml", '<p id="one">before <b>bold</b> after</p>')
    assert [fact.value for fact in actual] == ["one", "before ", "bold", " after"]
    assert actual[0].location == "/p[1]/@id"
    assert actual[2].location.startswith("/p[1]/b[1]/text()")
    assert stream.exhausted


def test_xml_entities_are_refused(tmp_path):
    from sage.source_adapters.base import SourceReadError

    with pytest.raises(SourceReadError):
        facts(tmp_path, ".xml", '<!DOCTYPE x [<!ENTITY e "expand">]><x>&e;</x>')


def test_xlsx_late_rows_have_stable_addresses_and_native_types(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["label", "value"])
    sheet.cell(20, 2, "late narrative")
    sheet.cell(21, 1, True)
    second = workbook.create_sheet("Other")
    second.cell(3, 1, "second sheet")
    workbook.save(path)
    stream = source_facts(path, None, FactLimits())
    actual = list(stream)
    assert [(fact.location, fact.value) for fact in actual][-3:] == [
        ("Data!B20", "late narrative"),
        ("Data!A21", True),
        ("Other!A3", "second sheet"),
    ]
    assert stream.exhausted


def test_narrative_projection_preserves_heading_and_preamble(tmp_path):
    projection = ProjectionResult(
        "all",
        [HeadingNode(1, "Topic", "Topic", "One.\n\nTwo.")],
        "hash",
        "1",
        "title",
        preamble="Opening.",
    )
    stream = source_facts(tmp_path / "source.md", projection, FactLimits())
    actual = list(stream)
    assert [fact.value for fact in actual] == ["Opening.", "One.", "Two."]
    assert [fact.region for fact in actual] == ["preamble", "Topic", "Topic"]
    assert actual[-1].context == "Topic"


@pytest.mark.parametrize("suffix", [".yaml", ".toml"])
def test_comments_preserved_but_hash_inside_string_is_not_comment(tmp_path, suffix):
    text = (
        'key: "# value" # authored comment\n'
        if suffix == ".yaml"
        else 'key = "# value" # authored comment\n'
    )
    actual, _ = facts(tmp_path, suffix, text)
    assert [fact.value for fact in actual] == ["# value", "# authored comment"]
    assert actual[-1].location == "line:1:comment"


def test_toml_comments_cannot_hide_depth_guard(tmp_path, monkeypatch):
    import tomllib

    calls = []
    monkeypatch.setattr(tomllib, "loads", lambda value: calls.append(value))
    actual, stream = facts(tmp_path, ".toml", "# ]]]]]\na = [[[[1]]]]", max_depth=2)
    assert actual == []
    assert calls == []
    assert stream.limitations


def test_xlsx_parser_guard_precedes_native_allocations(tmp_path, monkeypatch):
    import openpyxl

    path = tmp_path / "source.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["hello"])
    workbook.save(path)
    monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: pytest.fail("native load"))
    stream = source_facts(path, None, FactLimits(max_parse_nodes=2))
    assert list(stream) == []
    assert "before native allocation" in stream.limitations[0]


@pytest.mark.parametrize(
    ("name", "maximum"),
    [
        ("max_source_bytes", 256 * 1024 * 1024),
        ("max_depth", 128),
        ("max_parse_nodes", 1_600_000),
        ("max_rows", 400_000),
        ("max_columns", 1024),
    ],
)
def test_hard_fact_limits_refuse_allocation_expansion(name, maximum):
    with pytest.raises(ValueError, match="maximum"):
        FactLimits(**{name: maximum + 1})


def test_fact_stream_exposes_frozen_effective_limits(tmp_path):
    from dataclasses import FrozenInstanceError

    limits = FactLimits(max_rows=2)
    stream = source_facts(tmp_path / "source.md", None, limits)
    assert stream.limits is limits
    with pytest.raises(FrozenInstanceError):
        stream.limits.max_rows = 3


@pytest.mark.parametrize(
    "suffix,text",
    [
        (".json", '{"' + "k" * 600 + '": "value"}'),
        (".xml", "<" + "k" * 600 + ">value</" + "k" * 600 + ">"),
    ],
)
def test_location_guard_precedes_accumulated_paths(tmp_path, suffix, text):
    actual, stream = facts(tmp_path, suffix, text)
    assert actual == []
    assert "location" in stream.limitations[0]
    assert not stream.exhausted


def test_xml_comments_and_processing_instructions_are_authored(tmp_path):
    actual, _ = facts(tmp_path, ".xml", "<?note before?><p><!--explanation-->text</p>")
    assert [fact.value for fact in actual] == ["before", "explanation", "text"]


def test_numeric_native_map_preserves_actual_authored_labels(tmp_path):
    actual, _ = facts(tmp_path, ".json", '{"outer": {"useful label": 7}}')
    assert actual[0].authored_labels == ("outer", "useful label")
    assert actual[0].value == 7


def test_embedded_container_like_text_remains_unresolved(tmp_path):
    actual, stream = facts(tmp_path, ".json", '{"content": "[1,2]"}')
    assert [fact.value for fact in actual] == ["[1,2]"]
    assert any("embedded structure unresolved" in item for item in stream.limitations)


def test_long_narrative_paragraph_keeps_late_text_and_stable_offsets(tmp_path):
    text = "a" * 40000 + " late text"
    projection = ProjectionResult(text, [], "hash", "1", "title")
    stream = source_facts(tmp_path / "source.md", projection, FactLimits())
    actual = list(stream)
    assert [len(fact.value) for fact in actual] == [8192, 8192, 8192, 8192, 7242]
    assert "".join(fact.value for fact in actual) == text
    assert actual[-1].location == "document:char:32768"
    assert actual[-1].value.endswith(" late text")


@pytest.mark.parametrize(
    ("suffix", "text", "module_name", "materializer"),
    [
        (".json", "[[[[1]]]]", "json", "loads"),
        (".yaml", "a: [[[[1]]]]", "yaml", "safe_load_all"),
        (".toml", "a = [[[[1]]]]", "tomllib", "loads"),
    ],
)
def test_projection_guard_refuses_before_whole_materializer(
    tmp_path, monkeypatch, suffix, text, module_name, materializer
):
    import importlib

    from sage.source_adapters.base import SourceReadError
    from sage.source_adapters.structure import check_projection_work

    path = tmp_path / f"source{suffix}"
    path.write_text(text)
    calls = []
    module = importlib.import_module(module_name)
    monkeypatch.setattr(module, materializer, lambda value: calls.append(value))
    with pytest.raises(SourceReadError, match="projection work"):
        check_projection_work(path, FactLimits(max_depth=2))
        getattr(module, materializer)(text)
    assert calls == []


def test_projection_guard_aggregates_jsonl_parser_nodes(tmp_path, monkeypatch):
    import json

    from sage.source_adapters.base import SourceReadError
    from sage.source_adapters.structure import check_projection_work

    path = tmp_path / "source.jsonl"
    path.write_text("[[], []]\n[[], []]\n[[], []]\n")
    calls = []
    monkeypatch.setattr(json, "loads", lambda value: calls.append(value))
    with pytest.raises(SourceReadError, match="projection work"):
        check_projection_work(path, FactLimits(max_parse_nodes=6))
        json.loads("[]")
    assert calls == []


@pytest.mark.parametrize(
    "suffix,text", [(".csv", "a,b,c\n"), (".tsv", "a\tb\tc\n"), (".csv", "a\nb\nc\n")]
)
def test_projection_guard_refuses_delimited_work_before_projection(
    tmp_path, monkeypatch, suffix, text
):
    from sage.source_adapters import structured_data_adapter
    from sage.source_adapters.base import SourceReadError
    from sage.source_adapters.structure import check_projection_work

    path = tmp_path / f"source{suffix}"
    path.write_text(text)
    calls = []
    monkeypatch.setattr(structured_data_adapter, "_delimited_text", lambda *a: calls.append(a))
    with pytest.raises(SourceReadError, match="projection work"):
        check_projection_work(path, FactLimits(max_rows=2, max_columns=2))
        structured_data_adapter._delimited_text(text, ",")
    assert calls == []


def test_projection_guard_refuses_yaml_alias_before_construction(tmp_path, monkeypatch):
    import yaml

    from sage.source_adapters.base import SourceReadError
    from sage.source_adapters.structure import check_projection_work

    path = tmp_path / "source.yaml"
    path.write_text(
        "a: &a {key: value}\nb: &b {<<: [*a, *a]}\nc: &c {<<: [*b, *b]}\n"
        "d: &d {<<: [*c, *c]}\ne: {<<: [*d, *d]}\n"
    )
    monkeypatch.setattr(yaml, "safe_load_all", lambda *a: pytest.fail("native constructor reached"))
    with pytest.raises(SourceReadError, match="alias"):
        check_projection_work(path, FactLimits(max_parse_nodes=100))


def test_projection_guard_refuses_xml_before_tree_builder(tmp_path, monkeypatch):
    from sage.source_adapters import structured_data_adapter
    from sage.source_adapters.base import SourceReadError
    from sage.source_adapters.structure import check_projection_work

    path = tmp_path / "source.xml"
    path.write_text("<a><b><c>value</c></b></a>")
    monkeypatch.setattr(
        structured_data_adapter, "_xml_parse", lambda *a: pytest.fail("tree builder reached")
    )
    with pytest.raises(SourceReadError, match="projection work"):
        check_projection_work(path, FactLimits(max_depth=2))
        structured_data_adapter._xml_parse("")


def test_projection_guard_refuses_xlsx_before_workbook_allocation(tmp_path, monkeypatch):
    import openpyxl

    from sage.source_adapters.base import SourceReadError
    from sage.source_adapters.structure import check_projection_work

    path = tmp_path / "source.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.append(["value"])
    workbook.save(path)
    monkeypatch.setattr(
        openpyxl, "load_workbook", lambda *a, **k: pytest.fail("workbook allocated")
    )
    with pytest.raises(SourceReadError, match="projection work"):
        check_projection_work(path, FactLimits(max_parse_nodes=2))
        openpyxl.load_workbook(path)


def test_projection_guard_admits_before_reading(tmp_path, monkeypatch):
    from sage.source_adapters.base import SourceReadError
    from sage.source_adapters.structure import check_projection_work

    path = tmp_path / "source.json"
    path.write_text('["hello"]')
    monkeypatch.setattr(Path, "open", lambda *a, **k: pytest.fail("source opened"))
    with pytest.raises(SourceReadError, match="projection work"):
        check_projection_work(path, FactLimits(max_source_bytes=2))


@pytest.mark.asyncio
async def test_projection_guard_preserves_benign_yaml_alias_projection(tmp_path):
    from sage.source_adapters.structure import check_projection_work
    from sage.source_adapters.structured_data_adapter import StructuredDataAdapter

    path = tmp_path / "source.yaml"
    text = "a: &a {key: value}\nb: {<<: *a, other: text}\n"
    path.write_text(text)
    check_projection_work(path, FactLimits())
    projection = await StructuredDataAdapter().project(path)
    assert "*a" in projection.text
    assert "other: text" in projection.text


def test_projection_guard_refuses_recursive_yaml_alias(tmp_path):
    from sage.source_adapters.base import SourceReadError
    from sage.source_adapters.structure import check_projection_work

    path = tmp_path / "source.yaml"
    path.write_text("a: &a [*a]\n")
    with pytest.raises(SourceReadError, match="recursive/incomplete alias"):
        check_projection_work(path, FactLimits())


def test_xml_prolog_comments_and_instructions_have_source_facts(tmp_path):
    source = tmp_path / "prolog.xml"
    source.write_text("<!--intro--><?tool meta?><root>body</root><!--end-->")
    stream = source_facts(source, None, FactLimits())
    facts = list(stream)
    assert [f.rendering for f in facts] == ["intro", "meta", "body", "end"]
    assert facts[0].record == "document"
    assert facts[-1].record == "document"
    assert stream.exhausted


@pytest.mark.parametrize("suffix", ["json", "xlsx"])
async def test_native_adapter_entry_refuses_before_materializer(tmp_path, monkeypatch, suffix):
    import sage.source_adapters.structured_data_adapter as structured
    import sage.source_adapters.xlsx_adapter as xlsx
    from sage.source_adapters.structured_data_adapter import StructuredDataAdapter
    from sage.source_adapters.xlsx_adapter import XlsxAdapter

    path = tmp_path / ("source." + suffix)
    path.write_bytes(b"[]")
    adapter = StructuredDataAdapter() if suffix == "json" else XlsxAdapter()
    target = structured.json if suffix == "json" else xlsx
    monkeypatch.setattr(
        target,
        "loads" if suffix == "json" else "load_workbook",
        lambda *a, **k: pytest.fail("materializer reached"),
    )
    from sage.source_adapters.base import SourceReadError

    with pytest.raises(SourceReadError, match="source byte admission"):
        await adapter.project(path, {"discovery_facts": {"max_source_bytes": 1}})


async def test_native_json_entry_respects_explicit_nondefault_depth(tmp_path):
    from sage.source_adapters.base import SourceReadError
    from sage.source_adapters.structured_data_adapter import StructuredDataAdapter

    path = tmp_path / "deep.json"
    path.write_text("[" * 65 + '"text"' + "]" * 65)
    adapter = StructuredDataAdapter()
    with pytest.raises(SourceReadError, match="depth"):
        await adapter.project(path)
    projected = await adapter.project(path, {"discovery_facts": {"max_depth": 80}})
    import json

    assert json.loads(projected.text) == json.loads(path.read_text())


@pytest.mark.parametrize("unit,length,expected_chars", [("a", 32768, 8192), ("🙂", 32768, 2048)])
def test_narrative_default_utf8_segments_preserve_exact_divisible_source(
    tmp_path, unit, length, expected_chars
):
    text = unit * length
    projection = ProjectionResult(text, [], "hash", "1", "title")
    stream = source_facts(tmp_path / "source.md", projection, FactLimits())
    actual = list(stream)
    assert all(len(fact.value) == expected_chars for fact in actual)
    assert "".join(fact.value for fact in actual) == text
    assert [fact.location for fact in actual] == [
        f"document:char:{offset}" for offset in range(0, length, expected_chars)
    ]
    assert stream.narrative_input_bytes == 8192


def test_narrative_explicit_utf8_budget_includes_authored_context(tmp_path):
    text = "🙂" * 8
    projection = ProjectionResult(text, [HeadingNode(1, "題", "題", text)], "hash", "1", "title")
    stream = source_facts(
        tmp_path / "source.md", projection, FactLimits(), narrative_input_bytes=13
    )
    actual = list(stream)
    assert [fact.value for fact in actual] == ["🙂🙂"] * 4
    assert [fact.location for fact in actual] == [f"題:char:{i}" for i in (0, 2, 4, 6)]
    assert all(len(f"{fact.context}\n\n{fact.rendering}".encode()) == 13 for fact in actual)
    assert stream.narrative_input_bytes == 13


def test_narrative_unicode_unit_larger_than_budget_stays_whole(tmp_path):
    projection = ProjectionResult("🙂🙂", [], "hash", "1", "title")
    stream = source_facts(tmp_path / "source.md", projection, FactLimits(), narrative_input_bytes=1)
    actual = list(stream)
    assert [fact.value for fact in actual] == ["🙂", "🙂"]
    assert [fact.location for fact in actual] == ["document:char:0", "document:char:1"]
    assert stream.exhausted


def test_narrative_context_comes_from_heading_not_reserved_region_names(tmp_path):
    projection = ProjectionResult(
        "lead\nxy",
        [HeadingNode(1, "document", "document", "xy")],
        "hash",
        "1",
        "title",
        preamble="lead",
    )
    facts = list(
        source_facts(tmp_path / "source.md", projection, FactLimits(), narrative_input_bytes=11)
    )
    assert [fact.value for fact in facts] == ["lead", "x", "y"]
    assert [fact.context for fact in facts] == ["", "document", "document"]
    assert [fact.region for fact in facts] == ["preamble", "document", "document"]
