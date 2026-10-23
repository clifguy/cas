"""Processing identity uses one bounded summary before durable work."""

import hashlib
import json
from copy import deepcopy

import pytest
from docx import Document as WordDocument
from docx.enum.style import WD_STYLE_TYPE

from sage.config import VaultConfig
from sage.models.enums import SourceType
from sage.models.schemas import IngestRequest
from sage.source_adapters.docx_adapter import DocxAdapter
from tests.helpers.pipeline_wait import drain_abstraction_queue


def _supported_docx_configuration():
    return {"heading_style_map": {f"custom_style_{i:03d}_" + "x" * 40: 1 for i in range(100)}}


@pytest.mark.parametrize("configuration_source", ["request", "vault"])
async def test_supported_large_docx_configuration_ingests_and_recomputes(
    ingestion_service, tmp_vault_dir, monkeypatch, configuration_source
):
    import sage.services.ingestion as ingestion

    configuration = _supported_docx_configuration()
    assert len(json.dumps(configuration).encode()) == 6423
    ingestion_service._adapters[SourceType.DOCX] = DocxAdapter()
    captured = []
    select = ingestion.select_discovery

    def capture(*args, **kwargs):
        package = select(*args, **kwargs)
        captured.append(package)
        return package

    monkeypatch.setattr(ingestion, "select_discovery", capture)
    if configuration_source == "vault":
        payload = ingestion_service._config.model_dump()
        payload["adapter_defaults"] = {"docx": deepcopy(configuration)}
        ingestion_service._config = VaultConfig.model_validate(payload)
    source = tmp_vault_dir / "supported-configuration.docx"
    authored = WordDocument()
    first_style = next(iter(configuration["heading_style_map"]))
    authored.styles.add_style(first_style, WD_STYLE_TYPE.PARAGRAPH)
    authored.add_paragraph("Configured authored heading", style=first_style)
    authored.add_paragraph("Authored body remains represented.")
    authored.save(source)
    result = await ingestion_service.ingest(
        IngestRequest(
            source=str(source),
            source_type=SourceType.DOCX,
            config=configuration if configuration_source == "request" else None,
        )
    )
    initial = await ingestion_service._store.get_document(result.document.id)
    assert initial.pipeline_status.value == "abstraction_complete"
    retention = initial.processing_coverage.source_retention
    inspection = initial.processing_coverage.inspection
    assert retention.scope == inspection.scope == "complete"
    assert retention.identity.policy == "coverage-v1"
    assert inspection.identity.policy == "bounded-discovery:1"
    for identity in (retention.identity, inspection.identity):
        assert identity.configuration["effective_configuration_omitted"] is True
        assert identity.configuration["effective_configuration_sha256"].startswith("sha256:")
        assert len(json.dumps(identity.configuration).encode()) <= 4096
        assert len(identity.configuration) <= 32
    frozen = json.loads(captured[0].identity.configuration_json)
    for identity, full_configuration in (
        (retention.identity, configuration),
        (inspection.identity, frozen),
    ):
        expected_hash = (
            "sha256:"
            + hashlib.sha256(
                json.dumps(full_configuration, sort_keys=True, separators=(",", ":")).encode()
            ).hexdigest()
        )
        assert identity.configuration["effective_configuration_sha256"] == expected_hash
    for key, value in configuration["heading_style_map"].items():
        assert frozen["adapter"]["heading_style_map"][key] == value
    assert len(captured[0].identity.configuration_json.encode()) > 4096
    assert (
        "Configured authored heading"
        in await ingestion_service._content_store.get_heading_paths(initial.id)
    )
    await ingestion_service.recompute_pipeline(initial.id)
    await drain_abstraction_queue(ingestion_service)
    updated = await ingestion_service._store.get_document(initial.id)
    assert updated.processing_coverage.source_retention == retention
    assert updated.processing_coverage.inspection.generation > inspection.generation
    assert updated.processing_coverage.inspection.identity == inspection.identity


def _byte_sized_configuration(size):
    return {"value": "x" * (size - len(json.dumps({"value": ""}).encode()))}


@pytest.mark.parametrize(
    ("case", "configuration", "inline"),
    [
        ("4095-bytes", _byte_sized_configuration(4095), True),
        ("4096-bytes", _byte_sized_configuration(4096), True),
        ("4097-bytes", _byte_sized_configuration(4097), False),
        ("32-keys", {str(i): i for i in range(32)}, True),
        ("33-keys", {str(i): i for i in range(33)}, False),
        ("unicode-escaping", {"value": "🙂" * 400}, False),
        ("32-keys-unicode", {**{str(i): i for i in range(31)}, "value": "🙂" * 400}, False),
        ("compact-spacing", {"value": "x" * 4084}, False),
    ],
    ids=lambda item: item if isinstance(item, str) else None,
)
def test_configuration_summary_obeys_actual_model_limits(case, configuration, inline):
    from pydantic import ValidationError

    from sage.models.coverage import ProcessingIdentity, bounded_processing_configuration

    if case in {"unicode-escaping", "compact-spacing", "32-keys-unicode"}:
        assert (
            len(json.dumps(configuration, separators=(",", ":"), ensure_ascii=False).encode())
            <= 4096
        )
        assert len(json.dumps(configuration).encode()) > 4096
    if case == "32-keys-unicode":
        assert len(configuration) == 32
    if case.endswith("bytes"):
        assert len(json.dumps(configuration).encode()) == int(case.split("-")[0])
    if case.endswith("keys"):
        assert len(configuration) == int(case.split("-")[0])
    original = deepcopy(configuration)
    summary = bounded_processing_configuration(configuration)
    identity = ProcessingIdentity.from_effective_configuration(
        source_hash="sha256:" + "0" * 64,
        adapter="test:1",
        policy="retention:1",
        configuration=configuration,
        tokenization="unknown",
    )
    assert identity.configuration == summary
    assert len(json.dumps(summary).encode()) <= 4096
    assert len(summary) <= 32
    if inline:
        assert summary == original
    else:
        assert summary == {
            "effective_configuration_sha256": "sha256:"
            + hashlib.sha256(
                json.dumps(original, sort_keys=True, separators=(",", ":")).encode()
            ).hexdigest(),
            "effective_configuration_omitted": True,
        }
        with pytest.raises(ValidationError):
            ProcessingIdentity(
                source_hash="sha256:" + "0" * 64,
                adapter="test:1",
                policy="retention:1",
                configuration=original,
                tokenization="unknown",
            )
    configuration["changed"] = "later mutation"
    assert identity.configuration == summary
    assert "changed" not in summary


@pytest.mark.parametrize("unicode_value", [False, True], ids=["spacing", "unicode"])
def test_selection_summary_preserves_full_frozen_configuration(unicode_value):
    from dataclasses import asdict

    from sage.models.discovery import SelectionConfig, SelectionIdentity
    from sage.source_adapters.structure import FactLimits

    configuration = {
        "selection": asdict(SelectionConfig()),
        "facts": asdict(FactLimits()),
        "adapter": {"value": ""},
    }
    if unicode_value:
        configuration["adapter"]["value"] = "🙂" * 400
        frozen = json.dumps(
            configuration, sort_keys=True, separators=(",", ":"), ensure_ascii=False
        )
    else:
        initial = json.dumps(configuration, sort_keys=True, separators=(",", ":"))
        configuration["adapter"]["value"] = "x" * (4096 - len(initial.encode()))
        frozen = json.dumps(configuration, sort_keys=True, separators=(",", ":"))
    assert len(frozen.encode()) <= 4096
    assert len(json.dumps(configuration).encode()) > 4096
    original = deepcopy(configuration)
    captured = SelectionIdentity(
        "sha256:" + "0" * 64, "docx:1", "bounded-discovery:1", frozen, "utf8-byte-tokens:1"
    )
    configuration["adapter"]["value"] = "changed after capture"
    actual = captured.processing_identity()
    assert json.loads(captured.configuration_json) == original
    assert actual.configuration["effective_configuration_omitted"] is True
    assert (
        actual.configuration["effective_configuration_sha256"]
        == "sha256:"
        + hashlib.sha256(
            json.dumps(original, sort_keys=True, separators=(",", ":")).encode()
        ).hexdigest()
    )
    assert actual.configuration["selection"] == original["selection"]
    assert actual.configuration["facts"] == original["facts"]
    assert actual.configuration["adapter_configuration_omitted"] is True
    assert len(json.dumps(actual.configuration).encode()) <= 4096
    assert len(actual.configuration) <= 32


@pytest.mark.parametrize("operation", ["new", "force", "supersede"])
async def test_configuration_identity_admission_precedes_durable_work(
    ingestion_service, tmp_vault_dir, monkeypatch, operation
):
    from sage.api.errors import AdapterConfigInvalidError
    from sage.models.coverage import ProcessingIdentity

    source = tmp_vault_dir / "identity-admission.md"
    source.write_text("# Retained source\n\nAn authored paragraph.\n")
    predecessor = None
    if operation != "new":
        seeded = await ingestion_service.ingest(
            IngestRequest(source=str(source), source_type=SourceType.MARKDOWN)
        )
        predecessor = seeded.document.id
    if operation == "supersede":
        source = tmp_vault_dir / "identity-successor.md"
        source.write_text("# Successor source\n\nDifferent authored paragraph.\n")
    before = [
        item.model_dump(mode="json") for item in await ingestion_service._store.list_all_documents()
    ]
    retained_calls = []
    retain = ingestion_service._retain_or_reuse

    async def record_retention(*args, **kwargs):
        retained_calls.append(True)
        return await retain(*args, **kwargs)

    factory = ProcessingIdentity.from_effective_configuration

    def refuse(cls, **kwargs):
        if kwargs["policy"] == "coverage-v1":
            raise ValueError("controlled identity representation refusal")
        return factory(**kwargs)

    monkeypatch.setattr(ingestion_service, "_retain_or_reuse", record_retention)
    monkeypatch.setattr(ProcessingIdentity, "from_effective_configuration", classmethod(refuse))
    with pytest.raises((AdapterConfigInvalidError, ValueError)) as caught:
        await ingestion_service.ingest(
            IngestRequest(
                source=str(source),
                source_type=SourceType.MARKDOWN,
                force=operation == "force",
                predecessor_id=predecessor if operation == "supersede" else None,
            )
        )
    assert not retained_calls
    assert [
        item.model_dump(mode="json") for item in await ingestion_service._store.list_all_documents()
    ] == before
    assert isinstance(caught.value, AdapterConfigInvalidError)
