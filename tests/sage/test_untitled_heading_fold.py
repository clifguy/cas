"""A heading with no text shares no address with the text under no heading.

The empty heading path addresses the text before a document's first heading. A
heading whose text is empty is not a heading, so after ingest it is addressed by
nothing: the text under it is read in the section before it, and ``read_section``
with the empty path returns the text before the first heading and nothing else.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from sage.adapters.content_store_postgres import PostgresContentStore
from sage.adapters.stubs import StubAbstractionProvider, StubEmbeddingProvider
from sage.models.enums import SourceType
from sage.models.schemas import IngestRequest
from sage.services.ingestion import IngestionService
from sage.services.utilities import UtilitiesService
from sage.source_adapters.docx_adapter import DocxAdapter
from sage.source_adapters.markdown_adapter import MarkdownAdapter
from tests.helpers.pipeline_wait import await_pipeline_idle
from tests.sage.test_adapters import (
    _docx_untitled_heading_count,
    _markdown_untitled_heading_count,
    requires_docx,
)

pytestmark = pytest.mark.asyncio


class Indexed:
    def __init__(self, ingestion, store, utilities, graph_store, root):
        self.ingestion = ingestion
        self.store = store
        self.utilities = utilities
        self.graph_store = graph_store
        self.root = Path(root)

    async def ingest(self, path: Path, source_type: SourceType) -> str:
        result = await self.ingestion.ingest(
            IngestRequest(source=path.name, source_type=source_type)
        )
        await await_pipeline_idle(self.graph_store, result.document.id, service=self.ingestion)
        return result.document.id

    async def section(self, document_id: str, heading_path: str) -> str:
        return (await self.utilities.read_section(document_id, heading_path)).section_text


@pytest.fixture
async def indexed(
    graph_store, lock_manager, pg_pool, minimal_config, lifecycle_service, tmp_vault_dir
):
    store = PostgresContentStore(pg_pool)
    ingestion = IngestionService(
        graph_store=graph_store,
        lock_manager=lock_manager,
        content_store=store,
        embedding_provider=StubEmbeddingProvider(),
        abstraction_provider=StubAbstractionProvider(),
        config=minimal_config,
        source_adapters={SourceType.MARKDOWN: MarkdownAdapter(), SourceType.DOCX: DocxAdapter()},
        lifecycle_service=lifecycle_service,
    )
    utilities = UtilitiesService(
        graph_store=graph_store,
        content_store=store,
        embedding_provider=StubEmbeddingProvider(),
        config=minimal_config,
    )
    return Indexed(ingestion, store, utilities, graph_store, Path(tmp_vault_dir) / "sources")


def _markdown(indexed: Indexed, name: str, source: str) -> Path:
    assert _markdown_untitled_heading_count(source) == 1, "control: one untitled heading"
    path = indexed.root / name
    path.write_text(source)
    return path


async def test_an_untitled_heading_between_named_ones_is_read_in_the_section_before_it(indexed):
    path = _markdown(
        indexed,
        "between.md",
        "Lead sentinel.\n\n# Named\n\nNamed body.\n\n#\n\nOrphan body.\n\n# After\n\nAfter body.\n",
    )

    document_id = await indexed.ingest(path, SourceType.MARKDOWN)

    assert await indexed.store.get_heading_paths(document_id) == ["", "Named", "After"]
    assert await indexed.section(document_id, "") == "Lead sentinel."
    assert await indexed.section(document_id, "Named") == "# Named\n\nNamed body.\n\nOrphan body."


async def test_an_untitled_last_heading_leaves_the_empty_path_to_the_text_before_the_first(
    indexed,
):
    path = _markdown(
        indexed, "last.md", "Lead sentinel.\n\n# Goal\n\nGoal body.\n\n# Targets\n\n#\n\nTail.\n"
    )

    document_id = await indexed.ingest(path, SourceType.MARKDOWN)

    assert await indexed.store.get_heading_paths(document_id) == ["", "Goal", "Targets"]
    assert await indexed.section(document_id, "") == "Lead sentinel."
    assert await indexed.section(document_id, "Targets") == "# Targets\n\nTail."


@requires_docx
async def test_an_untitled_docx_heading_leaves_the_empty_path_to_the_text_before_the_first(
    indexed,
):
    import docx

    doc = docx.Document()
    doc.add_paragraph("Lead sentinel.")
    doc.add_paragraph("Goal", style="Heading 1")
    doc.add_paragraph("Goal body.")
    doc.add_paragraph("", style="Heading 1")
    doc.add_paragraph("Tail.")
    path = indexed.root / "last.docx"
    doc.save(str(path))
    assert _docx_untitled_heading_count(path) == 1, "control: one untitled heading"

    document_id = await indexed.ingest(path, SourceType.DOCX)

    assert await indexed.store.get_heading_paths(document_id) == ["", "Goal"]
    assert await indexed.section(document_id, "") == "Lead sentinel."
    assert await indexed.section(document_id, "Goal") == "# Goal\n\nGoal body.\nTail."
