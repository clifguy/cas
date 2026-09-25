"""Recovering the text documents carry before their first heading.

A document indexed before that text had a passage holds none of it, and its
stored passages cannot supply it. The migration re-projects the source, and when
the projection carries such text it adds that text as the document's first
passage. A stored passage the current adapter shapes the same way is kept as it
reads, so its heading path, section read and cached address resolve as before;
one an older adapter shaped differently is replaced by the current shape. Nothing
is re-abstracted.

Each document here is ingested through an adapter that withholds the preamble --
which is how the passage builder stored every document before -- and the
migration then runs with the adapter as shipped.
"""

from __future__ import annotations

import asyncio
from pathlib import Path

import pytest

from sage.adapters.content_store_postgres import PostgresContentStore
from sage.adapters.interfaces import Chunk
from sage.adapters.stubs import StubAbstractionProvider, StubEmbeddingProvider
from sage.models.enums import RetrievalMode, SourceType
from sage.models.schemas import DiscoverRequest, Document, IngestRequest, SetLifecycleRequest
from sage.services.ingestion import IngestionService
from sage.services.lifecycle import LifecycleService
from sage.services.maintenance import BACKFILL_TEXT_BEFORE_FIRST_HEADING, MaintenanceService
from sage.services.retrieval import RetrievalService
from sage.services.utilities import UtilitiesService
from sage.source_adapters.base import ProjectionResult, SourceAdapter
from sage.source_adapters.docx_adapter import DocxAdapter
from sage.source_adapters.markdown_adapter import MarkdownAdapter
from sage.source_adapters.pdf_adapter import PdfAdapter
from sage.storage.postgres.schema import EMBEDDING_DIM
from sage.vault_source_binding import FilesystemVaultSourceStore
from tests.helpers.pipeline_wait import await_pipeline_idle
from tests.sage.test_adapters import (
    _add_table,
    _make_pdf_with_outline,
    _make_pdf_with_pages,
    requires_docx,
    requires_pdf,
)

pytestmark = pytest.mark.asyncio

LEAD = "Larkspur branch holds the quarterly archive."
BODY = "# Guide\n\nGuide body.\n\n## Part\n\nPart body.\n"


class _RecordingEmbedder(StubEmbeddingProvider):
    """Records what it embeds, and can act between a read and its rewrite."""

    def __init__(self) -> None:
        super().__init__()
        self.embedded: list[str] = []
        self.during_embed = None

    async def embed(self, texts: list[str]) -> list[list[float]]:
        self.embedded.extend(texts)
        if self.during_embed is not None:
            await self.during_embed()
        return await super().embed(texts)


class _CountingStore(PostgresContentStore):
    """Counts every write of a document's passages."""

    def __init__(self, pool) -> None:
        super().__init__(pool)
        self.writes = 0

    async def index_chunks(self, document_id, chunks) -> None:
        self.writes += 1
        await super().index_chunks(document_id, chunks)


class _CountingAbstraction(StubAbstractionProvider):
    def __init__(self) -> None:
        self.calls = 0

    async def generate_abstract(self, text: str, max_tokens: int, doc_type: str | None) -> str:
        self.calls += 1
        return await super().generate_abstract(text, max_tokens, doc_type)


# The version each seeding adapter reports: the last before its adapter began
# reporting text before the first heading, which documents indexed then carry.
_EARLIER_VERSION = {MarkdownAdapter: "0.5.0", DocxAdapter: "0.4.0", PdfAdapter: "0.5.0"}


def _observed(base: type, *, withhold: bool):
    """An adapter of ``base`` that records each source it projects.

    Withholding the preamble, it also reports the version that preceded it, so the
    document it indexes reads as one an earlier adapter projected.
    """

    class Observed(base):
        VERSION = _EARLIER_VERSION[base] if withhold else base.VERSION

        def __init__(self) -> None:
            super().__init__()
            self.projected: list[str] = []

        async def project(self, source_path: Path, config: dict | None = None) -> ProjectionResult:
            projection = await super().project(source_path, config)
            self.projected.append(source_path.name)
            if withhold:
                projection.preamble = ""
            return projection

    return Observed()


class Vault:
    def __init__(self, graph_store, store, embedder, abstraction, ingestion, config, root):
        self.graph_store = graph_store
        self.store = store
        self.embedder = embedder
        self.abstraction = abstraction
        self.ingestion = ingestion
        self.config = config
        self.root = Path(root)

    async def ingest(self, name: str, content: bytes, source_type: SourceType) -> str:
        (self.root / "sources" / name).write_bytes(content)
        result = await self.ingestion.ingest(IngestRequest(source=name, source_type=source_type))
        await await_pipeline_idle(self.graph_store, result.document.id, service=self.ingestion)
        return result.document.id

    def ship_adapters(self) -> dict[SourceType, object]:
        """Swap in the adapters as shipped, each recording what it projects."""
        adapters = {
            SourceType.MARKDOWN: _observed(MarkdownAdapter, withhold=False),
            SourceType.DOCX: _observed(DocxAdapter, withhold=False),
            SourceType.PDF: _observed(PdfAdapter, withhold=False),
        }
        self.ingestion._adapters = adapters
        self.embedder.embedded.clear()
        self.store.writes = 0
        self.abstraction.calls = 0
        return adapters

    async def migrate(self):
        return await MaintenanceService(
            vault_id=self.config.vault.id,
            graph_store=self.graph_store,
            config=self.config,
            registry_service=None,
            content_store=self.store,
            ingestion_service=self.ingestion,
            vault_dir=self.root,
        ).migrate_vault()

    def utilities(self) -> UtilitiesService:
        return UtilitiesService(
            graph_store=self.graph_store,
            content_store=self.store,
            embedding_provider=StubEmbeddingProvider(),
            config=self.config,
        )


@pytest.fixture
async def vault(
    graph_store, lock_manager, pg_pool, minimal_config, lifecycle_service, tmp_vault_dir
):
    store = _CountingStore(pg_pool)
    embedder = _RecordingEmbedder()
    abstraction = _CountingAbstraction()
    ingestion = IngestionService(
        graph_store=graph_store,
        lock_manager=lock_manager,
        content_store=store,
        embedding_provider=embedder,
        abstraction_provider=abstraction,
        config=minimal_config,
        source_adapters={
            SourceType.MARKDOWN: _observed(MarkdownAdapter, withhold=True),
            SourceType.DOCX: _observed(DocxAdapter, withhold=True),
            SourceType.PDF: _observed(PdfAdapter, withhold=True),
        },
        lifecycle_service=lifecycle_service,
    )
    return Vault(
        graph_store, store, embedder, abstraction, ingestion, minimal_config, tmp_vault_dir
    )


async def _led(vault: Vault) -> str:
    return await vault.ingest("led.md", f"{LEAD}\n\n{BODY}".encode(), SourceType.MARKDOWN)


def _shape(chunks: list[Chunk]) -> list[tuple[str, str, int | None]]:
    return [(c.heading_path, c.content, c.section_index) for c in chunks]


async def test_the_text_before_the_first_heading_becomes_the_first_passage(vault):
    led = await _led(vault)
    before = await vault.store.get_all_chunks(led)
    before_paths = await vault.store.get_heading_paths(led)
    utilities = vault.utilities()
    before_sections = {p: (await utilities.read_section(led, p)).section_text for p in before_paths}
    before_abstract = (await vault.graph_store.get_document(led)).semantic_abstract
    assert "" not in before_paths, "control: the seeded document must lack the passage"
    assert before_abstract, "control: the document must carry an abstract to preserve"
    vault.ship_adapters()

    report = await vault.migrate()

    assert BACKFILL_TEXT_BEFORE_FIRST_HEADING in report.backfills_applied
    after = await vault.store.get_all_chunks(led)
    assert (after[0].heading_path, after[0].content, after[0].section_index) == ("", LEAD, 0)
    assert [(c.heading_path, c.content) for c in after[1:]] == [
        (c.heading_path, c.content) for c in before
    ]
    assert await vault.store.get_heading_paths(led) == ["", *before_paths]
    for path, text in before_sections.items():
        assert (await utilities.read_section(led, path)).section_text == text
    assert (await utilities.read_section(led, "")).section_text == LEAD
    assert (await vault.graph_store.get_document(led)).semantic_abstract == before_abstract
    assert vault.abstraction.calls == 0


async def test_a_second_run_reads_no_source_and_rewrites_nothing(vault):
    await _led(vault)
    await vault.ingest("plain.md", b"# Plain\n\nPlain body.\n", SourceType.MARKDOWN)
    first = vault.ship_adapters()
    await vault.migrate()
    assert vault.store.writes == 1, "control: the first run must have rewritten"
    assert sorted(first[SourceType.MARKDOWN].projected) == ["led.md", "plain.md"], (
        "control: the first run must have read both sources"
    )
    second = vault.ship_adapters()

    report = await vault.migrate()

    assert BACKFILL_TEXT_BEFORE_FIRST_HEADING not in report.backfills_applied
    assert second[SourceType.MARKDOWN].projected == []
    assert vault.embedder.embedded == []
    assert vault.store.writes == 0


async def test_an_examined_document_is_stamped_with_the_shipped_adapter_version(vault):
    led = await _led(vault)
    plain = await vault.ingest("plain.md", b"# Plain\n\nPlain body.\n", SourceType.MARKDOWN)
    for document_id in (led, plain):
        doc = await vault.graph_store.get_document(document_id)
        assert doc.adapter_version == "0.5.0", "control: seeded as an earlier adapter"
    vault.ship_adapters()

    await vault.migrate()

    for document_id in (led, plain):
        doc = await vault.graph_store.get_document(document_id)
        assert doc.adapter_version == MarkdownAdapter.VERSION


async def test_a_document_the_shipped_adapter_projected_is_not_read(vault):
    shipped = vault.ship_adapters()
    # Opening at a heading, it holds no empty-path passage, so only its version
    # can keep it from being read.
    await vault.ingest("fresh.md", BODY.encode(), SourceType.MARKDOWN)
    assert shipped[SourceType.MARKDOWN].projected == ["fresh.md"], "control: ingest projected it"
    later = vault.ship_adapters()

    report = await vault.migrate()

    assert later[SourceType.MARKDOWN].projected == []
    assert BACKFILL_TEXT_BEFORE_FIRST_HEADING not in report.backfills_applied


async def test_a_document_opening_at_a_heading_is_read_and_left_alone(vault):
    await _led(vault)
    plain = await vault.ingest(
        "plain.md", b"# Plain\n\nPlain sentinel body.\n", SourceType.MARKDOWN
    )
    before = _shape(await vault.store.get_all_chunks(plain))
    adapters = vault.ship_adapters()

    await vault.migrate()

    assert "plain.md" in adapters[SourceType.MARKDOWN].projected
    assert _shape(await vault.store.get_all_chunks(plain)) == before
    assert not any("Plain sentinel" in text for text in vault.embedder.embedded)
    assert any(LEAD in text for text in vault.embedder.embedded), (
        "control: the affected document in the same run was re-embedded"
    )


async def test_a_headingless_document_is_not_projected(vault):
    flat = await vault.ingest(
        "flat.md", b"Only a paragraph.\n\nAnd another.\n", SourceType.MARKDOWN
    )
    await _led(vault)
    before = _shape(await vault.store.get_all_chunks(flat))
    assert [path for path, _, _ in before] == [""], "control: the flat document's one passage"
    adapters = vault.ship_adapters()

    await vault.migrate()

    projected = adapters[SourceType.MARKDOWN].projected
    assert "led.md" in projected, "control: the migration projected its candidates"
    assert "flat.md" not in projected
    assert _shape(await vault.store.get_all_chunks(flat)) == before


@requires_pdf
async def test_every_earlier_pdf_is_examined_once_whatever_its_tags(vault, tmp_path):
    flat = await vault.ingest(
        "flat.pdf",
        _make_pdf_with_pages(
            tmp_path / "flat.pdf", [["PAGE_ONE"], ["PAGE_TWO"]], "Flat"
        ).read_bytes(),
        SourceType.PDF,
    )
    led = await vault.ingest(
        "led.pdf",
        _make_pdf_with_outline(
            tmp_path / "led.pdf",
            outline=[(1, "Intro", 1)],
            pages=[["COVER_PAGE_DELTA"], ["INTRO_BODY"]],
        ).read_bytes(),
        SourceType.PDF,
    )
    tags = (await vault.graph_store.get_document(led)).tags
    assert "pdf:has_outline" in tags, "control: the outline tag must be there to remove"
    await vault.graph_store.update_document(
        led, {"tags": [tag for tag in tags if tag != "pdf:has_outline"]}
    )
    before = _shape(await vault.store.get_all_chunks(flat))
    adapters = vault.ship_adapters()

    report = await vault.migrate()

    assert BACKFILL_TEXT_BEFORE_FIRST_HEADING in report.backfills_applied
    first = (await vault.store.get_all_chunks(led))[0]
    assert (first.heading_path, first.content) == ("", "COVER_PAGE_DELTA")
    assert sorted(adapters[SourceType.PDF].projected) == ["flat.pdf", "led.pdf"]
    assert _shape(await vault.store.get_all_chunks(flat)) == before
    again = vault.ship_adapters()
    await vault.migrate()
    assert again[SourceType.PDF].projected == []


async def test_a_source_changed_since_it_was_indexed_is_skipped_and_logged(vault, caplog):
    led = await _led(vault)
    doc = await vault.graph_store.get_document(led)
    (vault.root / "sources" / doc.source_path).write_text(f"Edited lead.\n\n{BODY}")
    before = _shape(await vault.store.get_all_chunks(led))
    vault.ship_adapters()

    with caplog.at_level("INFO", logger="sage.services.ingestion"):
        report = await vault.migrate()

    assert BACKFILL_TEXT_BEFORE_FIRST_HEADING not in report.backfills_applied
    assert _shape(await vault.store.get_all_chunks(led)) == before
    assert vault.embedder.embedded == []
    assert (await vault.graph_store.get_document(led)).adapter_version == "0.5.0", (
        "a skipped document is not stamped, so a later run examines it again"
    )
    skipped = [
        r.getMessage() for r in caplog.records if "not stored by migrate_vault" in r.getMessage()
    ]
    assert any(led in message and "source differs" in message for message in skipped), skipped


class _InMemorySourceStore(FilesystemVaultSourceStore):
    """Serves retained sources from memory, as a non-filesystem binding does."""

    def __init__(self, sources: dict[str, bytes]) -> None:
        super().__init__(Path("/unused/vault_root"))
        self.sources = sources
        self.reads: list[str] = []

    def source_exists(self, vault_id, storage_root, source_path):
        return source_path in self.sources

    def read_source(self, vault_id, storage_root, source_path):
        self.reads.append(source_path)
        return self.sources[source_path]


async def test_a_source_held_only_by_the_source_store_is_recovered(vault, monkeypatch):
    from sage.services.vault_source_errors import wrap_vault_source_store

    led = await _led(vault)
    plain = await vault.ingest("plain.md", b"# Plain\n\nPlain body.\n", SourceType.MARKDOWN)
    sources = {}
    for document_id in (led, plain):
        doc = await vault.graph_store.get_document(document_id)
        local = vault.root / "sources" / doc.source_path
        sources[doc.source_path] = local.read_bytes()
        local.unlink()
        assert not local.exists()
    remote = _InMemorySourceStore(sources)
    resolutions = []

    def resolve(*args, **kwargs):
        resolutions.append(1)
        return wrap_vault_source_store(remote)

    monkeypatch.setattr("sage.mcp_init.resolve_stack_vault_source_store", resolve)
    vault.ship_adapters()

    report = await vault.migrate()

    assert sorted(remote.reads) == sorted(sources)
    assert len(resolutions) == 1, "the source store is resolved once per run, not per document"
    assert BACKFILL_TEXT_BEFORE_FIRST_HEADING in report.backfills_applied
    assert (await vault.store.get_all_chunks(led))[0].content == LEAD


async def test_the_recovered_passage_carries_the_document_scalars(vault):
    led = await _led(vault)
    doc = await vault.graph_store.get_document(led)
    assert doc.doc_type and doc.lifecycle_status, "control: scalars to carry"
    vault.ship_adapters()

    await vault.migrate()

    first = (await vault.store.get_all_chunks(led))[0]
    assert first.heading_path == ""
    assert first.indexed_structure == ""
    assert (first.doc_type, first.lifecycle_status, first.project) == (
        doc.doc_type,
        str(doc.lifecycle_status),
        doc.project,
    )
    retrieval = RetrievalService(
        graph_store=vault.graph_store,
        content_store=vault.store,
        embedding_provider=StubEmbeddingProvider(),
        config=vault.config,
    )
    response = await retrieval.discover(
        DiscoverRequest(
            mode=RetrievalMode.KEYWORD,
            query="Larkspur",
            filters={"doc_type": doc.doc_type},
            limit=10,
        )
    )
    assert [h.document.id for h in response.results] == [led]


async def test_passages_already_holding_the_text_when_read_get_no_second(vault):
    """Candidacy is read from heading enumeration before the source is projected.

    The rewrite reads the passages again under the document's lock and does not
    trust that earlier read: passages already holding the empty path are left
    alone. A stale candidacy read is reproduced by answering heading enumeration
    with the paths the document held before its passages gained the text.
    """
    led = await _led(vault)
    stale_paths = await vault.store.get_heading_paths(led)
    assert "" not in stale_paths, "control: the seeded document must lack the passage"
    doc = await vault.graph_store.get_document(led)
    projection = await MarkdownAdapter().project(vault.root / "sources" / doc.source_path)
    reindexed = vault.ingestion._chunk_projection(led, projection)
    for chunk in reindexed:
        chunk.embedding = [0.25] * EMBEDDING_DIM
    await vault.store.index_chunks(led, reindexed)
    before = _shape(await vault.store.get_all_chunks(led))
    assert [path for path, _, _ in before].count("") == 1, "control: the re-index wrote it"

    real_paths = vault.store.get_heading_paths

    async def paths_as_first_read(document_id):
        return stale_paths if document_id == led else await real_paths(document_id)

    vault.store.get_heading_paths = paths_as_first_read
    vault.ship_adapters()

    report = await vault.migrate()

    assert BACKFILL_TEXT_BEFORE_FIRST_HEADING not in report.backfills_applied
    assert _shape(await vault.store.get_all_chunks(led)) == before
    assert vault.store.writes == 0


async def test_a_lifecycle_stamp_while_the_text_is_added_is_not_reverted(vault):
    """Archiving writes the new status onto the passage rows, which the rewrite
    carries forward from the rows it read. The archive waits for the rewrite."""
    led = await _led(vault)
    lifecycle = LifecycleService(
        vault.graph_store, vault.ingestion._locks, vault.config, vault.store
    )
    archiving = []

    async def archive_meanwhile():
        vault.embedder.during_embed = None
        archive = asyncio.create_task(
            lifecycle._set_lifecycle(led, SetLifecycleRequest(action="archive"))
        )
        archiving.append(archive)
        await asyncio.wait({archive}, timeout=0.5)
        assert not archive.done(), "control: the archive must wait on the rewrite"

    vault.ship_adapters()
    vault.embedder.during_embed = archive_meanwhile

    report = await vault.migrate()
    await archiving[0]

    assert BACKFILL_TEXT_BEFORE_FIRST_HEADING in report.backfills_applied
    chunks = await vault.store.get_all_chunks(led)
    assert chunks[0].heading_path == ""
    assert {c.lifecycle_status for c in chunks} == {"archived"}


async def test_a_document_with_no_adapter_to_project_it_is_skipped_and_logged(vault, caplog):
    led = await _led(vault)
    adapters = vault.ship_adapters()
    del adapters[SourceType.MARKDOWN]

    with caplog.at_level("INFO", logger="sage.services.ingestion"):
        report = await vault.migrate()

    assert BACKFILL_TEXT_BEFORE_FIRST_HEADING not in report.backfills_applied
    messages = [record.getMessage() for record in caplog.records]
    assert any(
        led in message and "not stored by migrate_vault" in message and "adapter" in message
        for message in messages
    ), messages


async def test_passages_an_older_adapter_shaped_differently_are_rebuilt(vault):
    """An older markdown adapter read a front-matter block's closing rule as a
    setext heading, filing the paragraph after it under that heading. The shipped
    adapter reports the paragraph as the text before the first heading, so adding
    it ahead of the stored passages would store it twice. The passages are rebuilt
    from the fresh projection instead, which is also what makes the version stamp
    true of them."""
    source = "---\ntitle: Probe\nstatus: active\n---\n\nIntro sentinel paragraph.\n\n" + BODY
    led = await vault.ingest("front.md", source.encode(), SourceType.MARKDOWN)
    stale = [
        Chunk(
            document_id=led,
            heading_path="title: Probe\nstatus: active",
            content="## title: Probe\nstatus: active\n\nIntro sentinel paragraph.",
            embedding=[0.25] * EMBEDDING_DIM,
            chunk_index=0,
            section_index=0,
        ),
        *(
            Chunk(
                document_id=led,
                heading_path=path,
                content=content,
                embedding=[0.25] * EMBEDDING_DIM,
                chunk_index=index,
                section_index=index,
            )
            for index, (path, content) in enumerate(
                [("Guide", "# Guide\n\nGuide body."), ("Guide > Part", "## Part\n\nPart body.")],
                start=1,
            )
        ),
    ]
    await vault.store.index_chunks(led, stale)
    await vault.graph_store.update_document(led, {"adapter_version": "0.4.0"})
    doc = await vault.graph_store.get_document(led)
    expected = vault.ingestion._chunk_projection(
        led, await MarkdownAdapter().project(vault.root / "sources" / doc.source_path)
    )
    vault.ship_adapters()

    report = await vault.migrate()

    assert BACKFILL_TEXT_BEFORE_FIRST_HEADING in report.backfills_applied
    after = await vault.store.get_all_chunks(led)
    assert [(c.heading_path, c.content) for c in after] == [
        (c.heading_path, c.content) for c in expected
    ]
    assert sum(c.content.count("Intro sentinel") for c in after) == 1
    assert await vault.store.get_heading_paths(led) == ["", "Guide", "Guide > Part"]
    assert (await vault.graph_store.get_document(led)).adapter_version == MarkdownAdapter.VERSION


@requires_docx
async def test_a_docx_document_an_earlier_adapter_projected_is_recovered(vault, tmp_path):
    import docx

    built = docx.Document()
    built.add_paragraph("Handbook Title", style="Title")
    built.add_paragraph("Opening paragraph gamma.")
    _add_table(built, [["Site", "Room"], ["Larkspur", "12"]])
    built.add_paragraph("Overview", style="Heading 1")
    built.add_paragraph("Overview body.")
    path = tmp_path / "lead.docx"
    built.save(str(path))
    led = await vault.ingest("lead.docx", path.read_bytes(), SourceType.DOCX)
    before = await vault.store.get_all_chunks(led)
    assert "" not in [c.heading_path for c in before], "control: seeded without the passage"
    assert (await vault.graph_store.get_document(led)).adapter_version == "0.4.0"
    adapters = vault.ship_adapters()

    report = await vault.migrate()

    assert adapters[SourceType.DOCX].projected == ["lead.docx"]
    assert BACKFILL_TEXT_BEFORE_FIRST_HEADING in report.backfills_applied
    first = (await vault.store.get_all_chunks(led))[0]
    assert first.heading_path == ""
    assert "gamma" in first.content and "Larkspur" in first.content
    assert (await vault.graph_store.get_document(led)).adapter_version == DocxAdapter.VERSION


# ── Untitled headings ───────────────────────────────────────────────
#
# An adapter that read a heading with no text as a heading stored the text under it
# at the empty path, or at a path with an empty segment. Where that path is the empty
# one it cannot be told from the text before the first heading by path alone, so the
# documents below are seeded with the passages such an adapter wrote, stamped with
# its version. A document whose empty path is its text under no heading is left
# unread beside them.

# The last version of each adapter to read a heading with no text as a heading.
_UNTITLED_READ_AS_HEADING = {
    SourceType.MARKDOWN: "0.6.0",
    SourceType.DOCX: "0.5.0",
    SourceType.PDF: "0.6.0",
}


async def _seeded(
    vault: Vault,
    name: str,
    source: str | bytes,
    sections,
    version: str,
    source_type: SourceType = SourceType.MARKDOWN,
) -> str:
    """Ingest ``source``, then store ``sections`` as its passages, stamped ``version``."""
    content = source.encode() if isinstance(source, str) else source
    document_id = await vault.ingest(name, content, source_type)
    await vault.store.index_chunks(
        document_id,
        [
            Chunk(
                document_id=document_id,
                heading_path=path,
                content=content,
                embedding=[0.25] * EMBEDDING_DIM,
                chunk_index=index,
                section_index=index,
            )
            for index, (path, content) in enumerate(sections)
        ],
    )
    await vault.graph_store.update_document(document_id, {"adapter_version": version})
    return document_id


async def _shipped_passages(
    vault: Vault, document_id: str, adapter: SourceAdapter | None = None
) -> list[tuple[str, str]]:
    doc = await vault.graph_store.get_document(document_id)
    adapter = adapter or MarkdownAdapter()
    projection = await adapter.project(vault.root / "sources" / doc.source_path)
    return [
        (c.heading_path, c.content)
        for c in vault.ingestion._chunk_projection(document_id, projection)
    ]


async def _untitled_opening(vault: Vault) -> str:
    return await _seeded(
        vault,
        "opening.md",
        "#\n\nUnder sentinel.\n\n# Named\n\nNamed body.\n",
        [("", "# \n\nUnder sentinel."), ("Named", "# Named\n\nNamed body.")],
        _UNTITLED_READ_AS_HEADING[SourceType.MARKDOWN],
    )


async def test_an_untitled_heading_stored_before_the_text_before_the_first_heading_is_recovered(
    vault,
):
    led = await _seeded(
        vault,
        "goal.md",
        f"{LEAD}\n\n# Goal\n\nGoal body.\n\n# Targets\n\n#\n\nTail.\n",
        [("Goal", "# Goal\n\nGoal body."), ("Targets", "# Targets"), ("", "# \n\nTail.")],
        _EARLIER_VERSION[MarkdownAdapter],
    )
    assert "" in await vault.store.get_heading_paths(led), (
        "control: the untitled heading holds the empty path"
    )
    expected = await _shipped_passages(vault, led)
    adapters = vault.ship_adapters()

    report = await vault.migrate()

    assert adapters[SourceType.MARKDOWN].projected == ["goal.md"]
    assert BACKFILL_TEXT_BEFORE_FIRST_HEADING in report.backfills_applied
    assert [(c.heading_path, c.content) for c in await vault.store.get_all_chunks(led)] == expected
    assert (await vault.utilities().read_section(led, "")).section_text == LEAD
    assert (await vault.utilities().read_section(led, "Targets")).section_text == (
        "# Targets\n\nTail."
    )
    assert (await vault.graph_store.get_document(led)).adapter_version == MarkdownAdapter.VERSION


async def test_an_untitled_heading_stored_beside_the_text_before_the_first_heading_is_rebuilt(
    vault,
):
    led = await _seeded(
        vault,
        "beside.md",
        f"{LEAD}\n\n# Named\n\nNamed body.\n\n#\n\nOrphan body.\n",
        [("", LEAD), ("Named", "# Named\n\nNamed body."), ("", "# \n\nOrphan body.")],
        _UNTITLED_READ_AS_HEADING[SourceType.MARKDOWN],
    )
    assert "Orphan" in (await vault.utilities().read_section(led, "")).section_text, (
        "control: the empty path addresses the untitled heading's text too"
    )
    adapters = vault.ship_adapters()

    await vault.migrate()

    assert adapters[SourceType.MARKDOWN].projected == ["beside.md"]
    assert (await vault.utilities().read_section(led, "")).section_text == LEAD
    assert (await vault.utilities().read_section(led, "Named")).section_text == (
        "# Named\n\nNamed body.\n\nOrphan body."
    )


async def test_an_untitled_heading_opening_a_document_is_read(vault):
    opening = await _untitled_opening(vault)
    expected = await _shipped_passages(vault, opening)
    adapters = vault.ship_adapters()

    await vault.migrate()

    assert adapters[SourceType.MARKDOWN].projected == ["opening.md"]
    assert [
        (c.heading_path, c.content) for c in await vault.store.get_all_chunks(opening)
    ] == expected
    assert expected[0] == ("", "Under sentinel."), "control: the fresh passage differs"


async def test_a_path_with_an_empty_segment_is_read(vault):
    nested = await _seeded(
        vault,
        "nested.md",
        "# A\n\nA body.\n\n##\n\nNested sentinel.\n",
        [("A", "# A\n\nA body."), ("A > ", "## \n\nNested sentinel.")],
        _UNTITLED_READ_AS_HEADING[SourceType.MARKDOWN],
    )
    adapters = vault.ship_adapters()

    await vault.migrate()

    assert adapters[SourceType.MARKDOWN].projected == ["nested.md"]
    assert await vault.store.get_heading_paths(nested) == ["A"]


async def test_text_before_the_first_heading_stored_by_its_adapter_is_not_read(vault):
    await _untitled_opening(vault)
    source = f"{LEAD}\n\n# Named\n\nNamed body.\n"
    led = await vault.ingest("genuine.md", source.encode(), SourceType.MARKDOWN)
    fresh = await _shipped_passages(vault, led)
    assert fresh[0] == ("", LEAD), "control: the empty path holds the text before the heading"
    await vault.store.index_chunks(
        led,
        [
            Chunk(
                document_id=led,
                heading_path=path,
                content=content,
                embedding=[0.25] * EMBEDDING_DIM,
                chunk_index=index,
                section_index=index,
            )
            for index, (path, content) in enumerate(fresh)
        ],
    )
    await vault.graph_store.update_document(
        led, {"adapter_version": _UNTITLED_READ_AS_HEADING[SourceType.MARKDOWN]}
    )
    adapters = vault.ship_adapters()

    await vault.migrate()

    projected = adapters[SourceType.MARKDOWN].projected
    assert "opening.md" in projected, "control: the migration projected its candidates"
    assert "genuine.md" not in projected


@pytest.mark.parametrize(
    "body", ["Only a paragraph.\n\nAnd another.\n", "#1 priority stays plain.\n\nAnd another.\n"]
)
async def test_a_headingless_document_stamped_by_an_untitled_reading_adapter_is_not_read(
    vault, body
):
    await _untitled_opening(vault)
    flat = await _seeded(
        vault,
        "flat.md",
        body,
        [("", body.strip())],
        _UNTITLED_READ_AS_HEADING[SourceType.MARKDOWN],
    )
    adapters = vault.ship_adapters()

    await vault.migrate()

    projected = adapters[SourceType.MARKDOWN].projected
    assert "opening.md" in projected, "control: the migration projected its candidates"
    assert "flat.md" not in projected
    assert await vault.store.get_heading_paths(flat) == [""]


async def test_an_empty_path_beside_headings_from_before_the_text_before_them_is_read(vault):
    """An adapter that stored no text before the first heading can only have put a
    heading at the empty path, whatever that passage opens with."""
    led = await _seeded(
        vault,
        "unmarked.md",
        f"{LEAD}\n\n# Goal\n\nGoal body.\n\n#\n\nTail.\n",
        [("Goal", "# Goal\n\nGoal body."), ("", "Tail.")],
        _EARLIER_VERSION[MarkdownAdapter],
    )
    adapters = vault.ship_adapters()

    await vault.migrate()

    assert adapters[SourceType.MARKDOWN].projected == ["unmarked.md"]
    assert (await vault.utilities().read_section(led, "")).section_text == LEAD


@requires_docx
async def test_a_docx_untitled_heading_beside_the_text_before_the_first_heading_is_rebuilt(
    vault, tmp_path
):
    import docx

    built = docx.Document()
    built.add_paragraph("Lead sentinel.")
    built.add_paragraph("Goal", style="Heading 1")
    built.add_paragraph("Goal body.")
    built.add_paragraph("", style="Heading 1")
    built.add_paragraph("Tail.")
    path = tmp_path / "beside.docx"
    built.save(str(path))
    led = await _seeded(
        vault,
        "beside.docx",
        path.read_bytes(),
        [("", "Lead sentinel."), ("Goal", "# Goal\n\nGoal body."), ("", "# \n\nTail.")],
        _UNTITLED_READ_AS_HEADING[SourceType.DOCX],
        SourceType.DOCX,
    )
    expected = await _shipped_passages(vault, led, DocxAdapter())
    adapters = vault.ship_adapters()

    await vault.migrate()

    assert adapters[SourceType.DOCX].projected == ["beside.docx"]
    assert [(c.heading_path, c.content) for c in await vault.store.get_all_chunks(led)] == expected
    assert (await vault.utilities().read_section(led, "")).section_text == "Lead sentinel."


@requires_pdf
async def test_a_pdf_heading_titled_by_its_outline_entry_object_is_read(vault, tmp_path):
    path = _make_pdf_with_outline(
        tmp_path / "repr.pdf",
        outline=[(1, "Named", 0), (1, "", 1), (1, "After", 2)],
        pages=[["NAMED_BODY"], ["ORPHAN_BODY"], ["AFTER_BODY"]],
    )
    untitled = "{'/Title': '', '/Page': IndirectObject(4, 0, 4402), '/Type': '/Fit'}"
    led = await _seeded(
        vault,
        "repr.pdf",
        path.read_bytes(),
        [
            ("Named", "# Named\n\nNAMED_BODY"),
            (untitled, f"# {untitled}\n\nORPHAN_BODY"),
            ("After", "# After\n\nAFTER_BODY"),
        ],
        _UNTITLED_READ_AS_HEADING[SourceType.PDF],
        SourceType.PDF,
    )
    paths = await vault.store.get_heading_paths(led)
    assert "" not in paths and all(segment.strip() for p in paths for segment in p.split(" > ")), (
        "control: no empty path or empty segment marks the untitled entry"
    )
    expected = await _shipped_passages(vault, led, PdfAdapter())
    adapters = vault.ship_adapters()

    await vault.migrate()

    assert adapters[SourceType.PDF].projected == ["repr.pdf"]
    assert [(c.heading_path, c.content) for c in await vault.store.get_all_chunks(led)] == expected
    assert await vault.store.get_heading_paths(led) == ["Named", "After"]


# ── Faithful to the document as it stands ───────────────────────────


@requires_docx
async def test_the_backfill_preserves_structure_projected_under_a_request_config(vault, tmp_path):
    """A heading only the ingest request's config made one stays a heading: the
    source is re-projected with the config it was projected with, so the only
    change is the text before the first heading."""
    import docx

    request_config = {"heading_style_map": {"Subtitle": 1}}
    # A vault default the request does not set, so the config reaching the
    # adapter shows whether the recorded request config was merged over the
    # vault's defaults or passed through alone.
    vault.ingestion._config = vault.config.model_copy(
        update={"adapter_defaults": {"docx": {"heading_style_map": {"Title": 1}}}}
    )
    built = docx.Document()
    built.add_paragraph("Opening paragraph delta.")
    built.add_paragraph("SUBTITLED", style="Subtitle")
    built.add_paragraph("Subtitle body.")
    path = vault.root / "sources" / "styled.docx"
    built.save(str(path))
    result = await vault.ingestion.ingest(
        IngestRequest(source="styled.docx", source_type=SourceType.DOCX, config=request_config)
    )
    await await_pipeline_idle(vault.graph_store, result.document.id, service=vault.ingestion)
    led = result.document.id
    before = await vault.store.get_all_chunks(led)
    assert await vault.store.get_heading_paths(led) == ["SUBTITLED"], (
        "control: seeded with the request's heading and without the text before it"
    )
    defaults_only = vault.ingestion._chunk_projection(led, await DocxAdapter().project(path))
    assert [c.heading_path for c in defaults_only] != ["", "SUBTITLED"], (
        "control: without the request config the source has no such heading"
    )
    shipped = vault.ship_adapters()[SourceType.DOCX]
    project = shipped.project
    received: list[dict | None] = []

    async def recording(source_path, config=None):
        received.append(config)
        return await project(source_path, config)

    shipped.project = recording

    report = await vault.migrate()

    assert BACKFILL_TEXT_BEFORE_FIRST_HEADING in report.backfills_applied
    after = await vault.store.get_all_chunks(led)
    assert (after[0].heading_path, after[0].content) == ("", "Opening paragraph delta.")
    assert [(c.heading_path, c.content) for c in after[1:]] == [
        (c.heading_path, c.content) for c in before
    ]
    assert received == [{"heading_style_map": {"Title": 1, "Subtitle": 1}}]


async def test_a_title_edit_before_the_lock_is_not_overwritten_by_the_rewrite(vault):
    """The title is read under the lock the rewrite holds, so an edit landing
    while the source is projected -- before that lock -- is the one the
    rewritten passages' structure derives from."""
    from sage.models.schemas import UpdateMetadataRequest
    from sage.services.metadata import MetadataService
    from sage.services.passage_structure import indexed_structure

    led = await _led(vault)
    metadata = MetadataService(vault.graph_store, vault.ingestion._locks, vault.config, vault.store)
    candidacy_title = (await vault.graph_store.get_document(led)).title
    assert indexed_structure("Guide > Part", candidacy_title) != indexed_structure(
        "Guide > Part", "Renamed"
    ), "control: the title read at candidacy must derive a different structure"
    adapters = vault.ship_adapters()
    shipped = adapters[SourceType.MARKDOWN]
    project = shipped.project

    async def project_then_rename(source_path, config=None):
        projection = await project(source_path, config)
        await metadata._update_metadata(
            led, UpdateMetadataRequest(title="Renamed"), modified_by="t"
        )
        return projection

    shipped.project = project_then_rename

    await vault.migrate()

    assert vault.store.writes == 1, "control: the passages must have been rewritten"
    chunks = await vault.store.get_all_chunks(led)
    assert chunks[0].heading_path == ""
    assert [c.indexed_structure for c in chunks] == [
        indexed_structure(c.heading_path, "Renamed") for c in chunks
    ]


# ---------------------------------------------------------------------------
# A document the backfill cannot repair is recorded, and not read again
# ---------------------------------------------------------------------------

_SKIP_LINE = "not stored by migrate_vault"


def _skip_lines(caplog, document_id: str) -> list[str]:
    return [
        r.getMessage()
        for r in caplog.records
        if _SKIP_LINE in r.getMessage() and document_id in r.getMessage()
    ]


def _not_repaired(report) -> dict[str, str]:
    return {entry.document_id: entry.reason for entry in report.documents_not_repaired}


async def _change_source(vault: Vault, led: str) -> None:
    doc = await vault.graph_store.get_document(led)
    (vault.root / "sources" / doc.source_path).write_text(f"Edited lead.\n\n{BODY}")


async def _remove_source(vault: Vault, led: str) -> None:
    doc = await vault.graph_store.get_document(led)
    (vault.root / "sources" / doc.source_path).unlink()


def _withdraw_adapter(adapters: dict) -> None:
    del adapters[SourceType.MARKDOWN]


async def _forget_source_path(vault: Vault, led: str, monkeypatch) -> None:
    """Serve the document with no source path, which the durable store cannot hold.

    The column is ``NOT NULL`` and the model requires a string, so the branch is
    reached only through a record built without validation -- the shape a
    binding that did not enforce the column would hand back.
    """
    list_all = vault.graph_store.list_all_documents

    async def with_no_source_path():
        docs = await list_all()
        return [
            Document.model_construct(**{**dict(d), "source_path": None}) if d.id == led else d
            for d in docs
        ]

    monkeypatch.setattr(vault.graph_store, "list_all_documents", with_no_source_path)


@pytest.mark.parametrize(
    "reason",
    ["source differs", "source not projected", "no adapter", "no source path"],
)
async def test_an_unrepairable_document_is_not_read_again(vault, caplog, monkeypatch, reason):
    """Each reason the backfill cannot repair a document is recorded on the run
    that finds it, reported with its reason, and not examined again.

    Anti-coincidental-pass: the first run must report the document and log it.
    Without that control a document that was never a candidate would satisfy
    every second-run assertion.
    """
    led = await _led(vault)
    if reason == "source differs":
        await _change_source(vault, led)
    elif reason == "source not projected":
        await _remove_source(vault, led)
    elif reason == "no source path":
        await _forget_source_path(vault, led, monkeypatch)
    adapters = vault.ship_adapters()
    if reason == "no adapter":
        _withdraw_adapter(adapters)

    with caplog.at_level("INFO", logger="sage.services.ingestion"):
        first = await vault.migrate()

    assert reason in _not_repaired(first).get(led, ""), first.documents_not_repaired
    assert len(_skip_lines(caplog, led)) == 1
    assert (await vault.graph_store.get_document(led)).adapter_version == "0.5.0", (
        "an unrepaired document is not claimed current"
    )

    caplog.clear()
    again = vault.ship_adapters()
    if reason == "no adapter":
        _withdraw_adapter(again)
    with caplog.at_level("INFO", logger="sage.services.ingestion"):
        second = await vault.migrate()

    assert again.get(SourceType.MARKDOWN) is None or again[SourceType.MARKDOWN].projected == []
    assert _skip_lines(caplog, led) == []
    assert second.documents_not_repaired == []


async def test_a_restored_copy_is_repaired_by_the_next_run(vault, tmp_path):
    """Restoring the retained copy undoes the record, so the next run repairs it.

    Anti-coincidental-pass: the record is confirmed to exclude the document
    first, so the repair is the restore's doing and not a record that never held.
    """
    led = await _led(vault)
    doc = await vault.graph_store.get_document(led)
    original = (vault.root / "sources" / doc.source_path).read_bytes()
    await _change_source(vault, led)
    vault.ship_adapters()
    assert led in _not_repaired(await vault.migrate())
    excluded = vault.ship_adapters()
    await vault.migrate()
    assert excluded[SourceType.MARKDOWN].projected == [], "control: the record excludes it"

    delivered = tmp_path / "original.md"
    delivered.write_bytes(original)
    service = MaintenanceService(
        vault_id=vault.config.vault.id,
        graph_store=vault.graph_store,
        config=vault.config,
        registry_service=None,
        content_store=vault.store,
        ingestion_service=vault.ingestion,
        vault_dir=vault.root,
    )
    await service.restore_vault_source_file(source=str(delivered), document_id=led)

    adapters = vault.ship_adapters()
    report = await vault.migrate()

    assert adapters[SourceType.MARKDOWN].projected == [Path(doc.source_path).name]
    assert BACKFILL_TEXT_BEFORE_FIRST_HEADING in report.backfills_applied
    assert report.documents_not_repaired == []
    assert (await vault.store.get_all_chunks(led))[0].content == LEAD
    assert (await vault.graph_store.get_document(led)).adapter_version == MarkdownAdapter.VERSION


async def test_an_adapter_registered_later_re_examines_a_recorded_document(vault):
    """The record names the adapter that examined the document -- none, here --
    so an adapter that could examine it now is not turned away by it."""
    led = await _led(vault)
    _withdraw_adapter(vault.ship_adapters())
    assert led in _not_repaired(await vault.migrate())

    adapters = vault.ship_adapters()
    report = await vault.migrate()

    assert adapters[SourceType.MARKDOWN].projected == ["led.md"]
    assert BACKFILL_TEXT_BEFORE_FIRST_HEADING in report.backfills_applied
    assert (await vault.store.get_all_chunks(led))[0].content == LEAD
    assert led not in await vault.graph_store.reprojection_skips(), (
        "a document brought current keeps no record that it could not be"
    )


async def test_restoring_an_intact_copy_clears_a_projection_failure_record(vault, tmp_path):
    """A projection that failed for a reason outside the file -- a format
    dependency missing on the host, say -- is recorded against intact bytes.
    Restoring those bytes writes nothing, and still clears the record, so the
    next run tries again once the host can project them.

    Anti-coincidental-pass: the record is confirmed to exclude the document, and
    the restore is confirmed to have taken the no-write path, before the repair.
    """
    led = await _led(vault)
    doc = await vault.graph_store.get_document(led)
    original = (vault.root / "sources" / doc.source_path).read_bytes()
    failing = vault.ship_adapters()

    async def cannot_project(source_path, config=None):
        raise ValueError("a dependency this host lacks")

    failing[SourceType.MARKDOWN].project = cannot_project
    assert "source not projected" in _not_repaired(await vault.migrate()).get(led, "")
    excluded = vault.ship_adapters()
    await vault.migrate()
    assert excluded[SourceType.MARKDOWN].projected == [], "control: the record excludes it"

    delivered = tmp_path / "original.md"
    delivered.write_bytes(original)
    restore = await MaintenanceService(
        vault_id=vault.config.vault.id,
        graph_store=vault.graph_store,
        config=vault.config,
        registry_service=None,
        content_store=vault.store,
        ingestion_service=vault.ingestion,
        vault_dir=vault.root,
    ).restore_vault_source_file(source=str(delivered), document_id=led)

    assert restore.status == "already_intact", "control: the restore wrote nothing"
    assert led not in await vault.graph_store.reprojection_skips()
    adapters = vault.ship_adapters()
    report = await vault.migrate()
    assert adapters[SourceType.MARKDOWN].projected == ["led.md"]
    assert BACKFILL_TEXT_BEFORE_FIRST_HEADING in report.backfills_applied
    assert (await vault.store.get_all_chunks(led))[0].content == LEAD


class _UnavailableSourceStore(FilesystemVaultSourceStore):
    """A source store that declines every read as a transient condition."""

    def __init__(self) -> None:
        super().__init__(Path("/unused/vault_root"))

    def source_exists(self, vault_id, storage_root, source_path):
        return True

    def read_source(self, vault_id, storage_root, source_path):
        from sage.api.errors import VaultSourceStoreUnavailableError

        raise VaultSourceStoreUnavailableError(source_path, "read")


async def test_a_transient_store_outage_is_not_recorded(vault, monkeypatch):
    """A store that may serve the source on a later attempt does not make a
    document unrepairable, so the next run with the store healthy repairs it."""
    from sage.services.vault_source_errors import wrap_vault_source_store

    led = await _led(vault)
    doc = await vault.graph_store.get_document(led)
    local = vault.root / "sources" / doc.source_path
    healthy = _InMemorySourceStore({doc.source_path: local.read_bytes()})
    local.unlink()
    stores = [_UnavailableSourceStore(), healthy]
    monkeypatch.setattr(
        "sage.mcp_init.resolve_stack_vault_source_store",
        lambda *a, **k: wrap_vault_source_store(stores.pop(0)),
    )
    vault.ship_adapters()

    first = await vault.migrate()

    assert "source not projected" in _not_repaired(first).get(led, ""), (
        "control: the outage is still reported on the run it happens"
    )
    assert led not in await vault.graph_store.reprojection_skips()

    report = await vault.migrate()

    assert healthy.reads == [doc.source_path]
    assert BACKFILL_TEXT_BEFORE_FIRST_HEADING in report.backfills_applied
    assert (await vault.store.get_all_chunks(led))[0].content == LEAD


async def test_a_reprojection_skip_is_recorded_listed_cleared_and_removed(vault):
    """The graph store's record of an unrepaired document, end to end."""
    led = await _led(vault)
    other = await vault.ingest("other.md", b"# Other\n\nOther body.\n", SourceType.MARKDOWN)
    assert await vault.graph_store.reprojection_skips() == {}

    await vault.graph_store.record_reprojection_skip(led, "0.7.0", "why")
    await vault.graph_store.record_reprojection_skip(led, "0.8.0", "why")
    await vault.graph_store.record_reprojection_skip(other, None, "why not")
    assert await vault.graph_store.reprojection_skips() == {led: "0.8.0", other: None}, (
        "a later record replaces an earlier one"
    )

    await vault.graph_store.clear_reprojection_skip(led)
    assert await vault.graph_store.reprojection_skips() == {other: None}

    await vault.graph_store.remove_document(other)
    assert await vault.graph_store.reprojection_skips() == {}


@pytest.mark.parametrize(
    ("fault", "recorded"),
    [
        (PermissionError(13, "Permission denied"), False),
        (BlockingIOError(35, "Resource temporarily unavailable"), False),
        (FileNotFoundError(2, "No such file or directory"), True),
        (IsADirectoryError(21, "Is a directory"), True),
        (NotADirectoryError(20, "Not a directory"), True),
        (ValueError("could not open the file"), True),
    ],
    ids=["permission", "busy", "not-found", "is-a-directory", "not-a-directory", "wrapped"],
)
async def test_a_read_fault_of_the_moment_is_not_recorded(vault, fault, recorded):
    """A read fault that may not recur -- a permission, a busy file -- is reported
    but not recorded, whichever source binding raised it, so the next run tries
    again. A fault that is a property of the recorded path is recorded, and so is
    one that arrives re-raised as something other than an operating-system error.

    Anti-coincidental-pass: every case is first confirmed reported, so a fault
    that never reached the backfill cannot pass the recording assertion.
    """
    led = await _led(vault)
    failing = vault.ship_adapters()

    async def cannot_read(source_path, config=None):
        raise fault

    failing[SourceType.MARKDOWN].project = cannot_read
    report = await vault.migrate()

    assert "source not projected" in _not_repaired(report).get(led, ""), "control: reported"
    assert (led in await vault.graph_store.reprojection_skips()) is recorded
    adapters = vault.ship_adapters()
    await vault.migrate()
    assert adapters[SourceType.MARKDOWN].projected == ([] if recorded else ["led.md"])
