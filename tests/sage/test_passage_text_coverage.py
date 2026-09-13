"""Every word a source authors reaches the document's stored passages.

Passages are the durable form of a projection: section reads, projection text
and both retrieval arms read them and nothing else. An adapter's ``text`` is
built alongside its headings rather than from them, so it is an independent
record of what the source authored, and any of its words missing from the
stored passages is text no reader can reach.

Each case carries authored text before its first heading, the position the
passage builder once dropped.
"""

from __future__ import annotations

import re
from collections.abc import Callable
from pathlib import Path

import pytest

from sage.models.enums import SourceType
from sage.models.schemas import IngestRequest
from sage.services.ingestion import IngestionService
from sage.source_adapters.docx_adapter import DocxAdapter
from sage.source_adapters.markdown_adapter import MarkdownAdapter
from sage.source_adapters.pdf_adapter import PdfAdapter
from sage.source_adapters.pptx_adapter import PptxAdapter
from sage.source_adapters.xlsx_adapter import XlsxAdapter
from tests.helpers.pipeline_wait import await_pipeline_idle
from tests.sage.test_adapters import (
    _add_table,
    _build_template_fixture,
    _make_multisheet_xlsx,
    _make_pdf_with_outline,
    _make_pptx,
    requires_docx,
    requires_openpyxl,
    requires_pdf,
    requires_pptx,
)

pytestmark = pytest.mark.asyncio

_FRONT_MATTER = re.compile(r"\A---\n.*?\n---\n", re.DOTALL)


def _markdown(directory: Path) -> Path:
    path = directory / "lead.md"
    path.write_text(
        "Opening sentence alpha.\n\n"
        "| Site | Room |\n|---|---|\n| Larkspur | 12 |\n\n"
        "# Guide\n\nGuide body.\n\n## Part\n\nPart body.\n"
    )
    return path


def _markdown_front_matter(directory: Path) -> Path:
    path = directory / "front.md"
    path.write_text("---\nname: excluded\n---\n\nIntro beta.\n\n# Heading\n\nBody.\n")
    return path


def _docx(directory: Path) -> Path:
    import docx

    doc = docx.Document()
    doc.add_paragraph("Handbook Title", style="Title")
    doc.add_paragraph("Opening paragraph gamma.")
    _add_table(doc, [["Site", "Room"], ["Larkspur", "12"]])
    doc.add_paragraph("Overview", style="Heading 1")
    doc.add_paragraph("Overview body.")
    path = directory / "lead.docx"
    doc.save(str(path))
    return path


def _dotx(directory: Path) -> Path:
    # A template's synthesized description of its style surface is the text
    # the adapter places first, ahead of the template's own first heading.
    return _build_template_fixture(directory, "lead.dotx")


def _pdf(directory: Path) -> Path:
    return _make_pdf_with_outline(
        directory / "lead.pdf",
        outline=[(1, "Intro", 1), (2, "Detail", 2)],
        pages=[["Cover page delta", "Prepared by epsilon"], ["Intro body"], ["Detail body"]],
    )


def _pptx(directory: Path) -> Path:
    return _make_pptx(
        directory,
        [
            {"title": None, "body": ["UNTITLED_SLIDE_BODY"]},
            {"title": "Second", "body": ["SECOND_BODY"], "notes": "SPEAKER_NOTES"},
        ],
        filename="lead.pptx",
    )


def _xlsx(directory: Path) -> Path:
    return _make_multisheet_xlsx(
        directory, {"First": [["FIRST_CELL"]], "Second": [["SECOND_CELL"]]}, filename="lead.xlsx"
    )


# Each case names words its fixture places before the first heading. The words
# checked are drawn from the adapter's own text, so these are the control that the
# text still carries them: an adapter dropping them from its text and its headings
# alike would otherwise leave nothing missing to find.
CASES = [
    pytest.param(SourceType.MARKDOWN, _markdown, {"alpha", "larkspur"}, id="markdown"),
    pytest.param(SourceType.MARKDOWN, _markdown_front_matter, {"beta"}, id="markdown-front-matter"),
    pytest.param(
        SourceType.DOCX, _docx, {"handbook", "gamma", "larkspur"}, id="docx", marks=requires_docx
    ),
    pytest.param(SourceType.DOCX, _dotx, {"appendix"}, id="dotx", marks=requires_docx),
    pytest.param(SourceType.PDF, _pdf, {"delta", "epsilon"}, id="pdf-outline", marks=requires_pdf),
    pytest.param(SourceType.PPTX, _pptx, {"untitled_slide_body"}, id="pptx", marks=requires_pptx),
    pytest.param(SourceType.XLSX, _xlsx, {"first_cell"}, id="xlsx", marks=requires_openpyxl),
]

ADAPTERS = {
    SourceType.MARKDOWN: MarkdownAdapter(),
    SourceType.DOCX: DocxAdapter(),
    SourceType.PDF: PdfAdapter(),
    SourceType.PPTX: PptxAdapter(),
    SourceType.XLSX: XlsxAdapter(),
}


def _words(text: str) -> set[str]:
    return {word.casefold() for word in re.findall(r"\w+", text)}


@pytest.fixture
def ingestion(
    graph_store,
    lock_manager,
    stub_content_store,
    stub_embedding_provider,
    stub_abstraction_provider,
    minimal_config,
    lifecycle_service,
):
    return IngestionService(
        graph_store=graph_store,
        lock_manager=lock_manager,
        content_store=stub_content_store,
        embedding_provider=stub_embedding_provider,
        abstraction_provider=stub_abstraction_provider,
        config=minimal_config,
        source_adapters=ADAPTERS,
        lifecycle_service=lifecycle_service,
    )


@pytest.mark.parametrize(("source_type", "build", "leading"), CASES)
async def test_every_authored_word_reaches_a_stored_passage(
    ingestion,
    graph_store,
    stub_content_store,
    tmp_vault_dir,
    source_type: SourceType,
    build: Callable[[Path], Path],
    leading: set[str],
):
    sources = Path(tmp_vault_dir) / "sources"
    path = build(sources)
    projection = await ADAPTERS[source_type].project(path)
    authored = _words(_FRONT_MATTER.sub("", projection.text))
    assert leading <= authored, f"control: the adapter's text lost {sorted(leading - authored)}"

    result = await ingestion.ingest(IngestRequest(source=path.name, source_type=source_type))
    await await_pipeline_idle(graph_store, result.document.id, service=ingestion)
    stored = await stub_content_store.get_all_chunks(result.document.id)

    missing = authored - _words("\n".join(chunk.content for chunk in stored))
    assert not missing, f"authored words absent from every stored passage: {sorted(missing)}"
